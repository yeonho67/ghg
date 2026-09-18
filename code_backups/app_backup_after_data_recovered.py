import os
import shutil
import sqlite3
from datetime import datetime
from io import BytesIO

import pandas as pd
import plotly.express as px
import streamlit as st


DB_NAME = "ghg_manager.db"

START_YEAR = 2021
MAX_YEAR = 2100

SITES = {
    "예산": {"country": "한국", "corporation": "예산법인"},
    "과천": {"country": "한국", "corporation": "비매출사업장"},
    "안성": {"country": "한국", "corporation": "안성법인"},
    "멕시코": {"country": "멕시코", "corporation": "멕시코법인"},
    "베트남": {"country": "베트남", "corporation": "베트남법인"},
    "중국": {"country": "중국", "corporation": "중국법인"},
}

ENERGY_TYPES = ["전기", "경유", "휘발유", "펠릿", "LPG", "LNG"]
COMBUSTION_TYPES = ["고정연소", "이동연소"]
MONTH_COLUMNS = [
    "1월", "2월", "3월", "4월", "5월", "6월",
    "7월", "8월", "9월", "10월", "11월", "12월"
]
UNIT_OPTIONS = ["MWh", "kWh", "L", "kg", "ton", "m³", "Nm³"]

def get_usage_units_by_energy(energy_type):
    unit_map = {
        "전기": ["MWh"],
        "경유": ["L"],
        "휘발유": ["L"],
        "LNG": ["m³"],
        "LPG": ["kg", "L", "m³"],
        "펠릿": ["kg"],
    }

    return unit_map.get(energy_type, UNIT_OPTIONS)

CORPORATIONS = [
    "예산법인",
    "안성법인",
    "멕시코법인",
    "베트남법인",
    "중국법인",
]


def get_connection():
    return sqlite3.connect(DB_NAME)


def get_current_year():
    return datetime.now().year


def get_default_year():
    current_year = get_current_year()

    if current_year < START_YEAR:
        return START_YEAR

    return current_year


def init_db():
    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS facility_energy (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            site TEXT NOT NULL,
            facility TEXT NOT NULL,
            energy_type TEXT NOT NULL,
            usage_unit TEXT NOT NULL,
            combustion_type TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT DEFAULT '',
            UNIQUE(site, facility)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS activity_data (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            year INTEGER NOT NULL,
            month INTEGER NOT NULL,
            site TEXT NOT NULL,
            facility TEXT NOT NULL,
            country TEXT NOT NULL,
            corporation TEXT NOT NULL,
            energy_type TEXT NOT NULL,
            usage_amount REAL NOT NULL,
            usage_unit TEXT NOT NULL,
            scope TEXT NOT NULL,
            combustion_type TEXT NOT NULL,
            emission_amount REAL NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT DEFAULT '',
            UNIQUE(year, month, site, facility)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS emission_factors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            factor_year INTEGER NOT NULL,
            country TEXT NOT NULL,
            energy_type TEXT NOT NULL,
            combustion_type TEXT NOT NULL,
            net_calorific_value REAL DEFAULT 0,
            emission_factor REAL NOT NULL,
            unit TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT DEFAULT '',
            UNIQUE(factor_year, country, energy_type, combustion_type)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS monthly_revenue (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            year INTEGER NOT NULL,
            month INTEGER NOT NULL,
            corporation TEXT NOT NULL,
            revenue_okr REAL NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT DEFAULT '',
            UNIQUE(year, month, corporation)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS reduction_targets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            year INTEGER NOT NULL,
            corporation TEXT NOT NULL,
            base_emission REAL NOT NULL,
            reduction_rate REAL NOT NULL,
            target_emission REAL NOT NULL,
            memo TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT DEFAULT '',
            UNIQUE(year, corporation)
        )
        """
    )

    # 기존 DB 마이그레이션: emission_factors 테이블에 usage_unit 컬럼 추가
    cur.execute("PRAGMA table_info(emission_factors)")
    emission_factor_columns = [row[1] for row in cur.fetchall()]

    if "usage_unit" not in emission_factor_columns:
        cur.execute(
            """
            ALTER TABLE emission_factors
            ADD COLUMN usage_unit TEXT DEFAULT ''
            """
        )

    # 기존 배출계수 데이터에 usage_unit 기본값 채우기
    cur.execute(
        """
        UPDATE emission_factors
        SET usage_unit = CASE
            WHEN energy_type = '전기' THEN 'MWh'
            WHEN energy_type = '경유' THEN 'L'
            WHEN energy_type = '휘발유' THEN 'L'
            WHEN energy_type = '펠릿' THEN 'kg'
            WHEN energy_type = 'LNG' THEN 'm³'
            WHEN energy_type = 'LPG' THEN 'L'
            ELSE ''
        END
        WHERE usage_unit IS NULL OR TRIM(usage_unit) = ''
        """
    )
    
    conn.commit()
    conn.close()


def get_country(site):
    return SITES[site]["country"]


def get_corporation(site):
    return SITES[site]["corporation"]


def get_scope(energy_type):
    if energy_type == "전기":
        return "Scope 2"

    return "Scope 1"


def get_default_unit(site, energy_type):
    if energy_type == "전기":
        return "MWh"

    if energy_type in ["경유", "휘발유"]:
        return "L"

    if energy_type == "펠릿":
        return "kg"

    if energy_type == "LNG":
        return "m³"

    if energy_type == "LPG":
        if site == "예산":
            return "kg"
        if site == "멕시코":
            return "L"
        if site == "중국":
            return "m³"
        return "kg"

    return ""


def get_default_combustion_type(energy_type):
    if energy_type in ["전기", "LNG", "LPG", "펠릿"]:
        return "고정연소"

    if energy_type in ["경유", "휘발유"]:
        return "이동연소"

    return "고정연소"


def seed_default_factors():
    years = list(range(START_YEAR, get_current_year() + 1))
    countries = ["한국", "멕시코", "베트남", "중국"]

    sample_factors = {
        "전기": {"ncv": 0, "ef": 0.45941, "unit": "tCO2e/MWh 예시"},
        "경유": {"ncv": 35.2, "ef": 74100, "unit": "TJ/Gg, kgCO2/TJ 예시"},
        "휘발유": {"ncv": 30.3, "ef": 69300, "unit": "TJ/Gg, kgCO2/TJ 예시"},
        "펠릿": {"ncv": 17.0, "ef": 112000, "unit": "TJ/Gg, kgCO2/TJ 예시"},
        "LPG": {"ncv": 50.2, "ef": 63100, "unit": "단위별 실제 계수 확인 필요"},
        "LNG": {"ncv": 54.6, "ef": 56100, "unit": "TJ/Gg, kgCO2/TJ 예시"},
    }

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    for factor_year in years:
        for country in countries:
            for energy_type in ENERGY_TYPES:
                for combustion_type in COMBUSTION_TYPES:
                    if energy_type == "전기" and combustion_type == "이동연소":
                        continue

                    item = sample_factors[energy_type]

                    cur.execute(
                        """
                        INSERT OR IGNORE INTO emission_factors
                        (
                            factor_year,
                            country,
                            energy_type,
                            combustion_type,
                            net_calorific_value,
                            emission_factor,
                            unit,
                            created_at,
                            updated_at
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            factor_year,
                            country,
                            energy_type,
                            combustion_type,
                            item["ncv"],
                            item["ef"],
                            item["unit"],
                            now_text,
                            now_text,
                        ),
                    )

    conn.commit()
    conn.close()


def load_factors():
    conn = get_connection()
    df = pd.read_sql_query(
        """
        SELECT
            id,
            factor_year,
            country,
            energy_type,
            combustion_type,
            net_calorific_value,
            emission_factor,
            unit,
            created_at,
            updated_at
        FROM emission_factors
        ORDER BY factor_year, country, energy_type, combustion_type
        """,
        conn,
    )
    conn.close()
    return df


def get_factor(factor_year, country, energy_type, combustion_type):
    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        SELECT
            net_calorific_value,
            emission_factor
        FROM emission_factors
        WHERE factor_year = ?
          AND country = ?
          AND energy_type = ?
          AND combustion_type = ?
        """,
        (
            factor_year,
            country,
            energy_type,
            combustion_type,
        ),
    )

    row = cur.fetchone()
    conn.close()

    if row is None:
        return None, None

    return row[0], row[1]


def update_factor(
    factor_year,
    country,
    energy_type,
    combustion_type,
    ncv,
    ef,
    unit
):
    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        INSERT INTO emission_factors
        (
            factor_year,
            country,
            energy_type,
            combustion_type,
            net_calorific_value,
            emission_factor,
            unit,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(factor_year, country, energy_type, combustion_type)
        DO UPDATE SET
            net_calorific_value = excluded.net_calorific_value,
            emission_factor = excluded.emission_factor,
            unit = excluded.unit,
            updated_at = excluded.updated_at
        """,
        (
            factor_year,
            country,
            energy_type,
            combustion_type,
            ncv,
            ef,
            unit,
            now_text,
            now_text,
        ),
    )

    conn.commit()
    conn.close()


def calculate_emission(
    factor_year,
    site,
    energy_type,
    usage_amount,
    combustion_type
):
    country = get_country(site)
    ncv, ef = get_factor(
        factor_year=factor_year,
        country=country,
        energy_type=energy_type,
        combustion_type=combustion_type,
    )

    if ncv is None or ef is None:
        return 0

    if usage_amount is None:
        usage_amount = 0

    if energy_type == "전기":
        return usage_amount * ef

    return usage_amount * ncv * ef * (10 ** -9)


def load_facility_energy():
    conn = get_connection()
    df = pd.read_sql_query(
        """
        SELECT
            id,
            site,
            facility,
            energy_type,
            usage_unit,
            combustion_type,
            created_at,
            updated_at
        FROM facility_energy
        ORDER BY site, facility
        """,
        conn,
    )
    conn.close()
    return df


def add_facility_energy(
    site,
    facility,
    energy_type,
    usage_unit,
    combustion_type
):
    facility = facility.strip()

    if facility == "":
        raise ValueError("시설명을 입력하세요.")

    if site not in SITES:
        raise ValueError("등록되지 않은 사업장입니다.")

    if energy_type not in ENERGY_TYPES:
        raise ValueError("등록되지 않은 에너지원입니다.")

    if usage_unit.strip() == "":
        raise ValueError("사용량 단위를 입력하세요.")

    if combustion_type not in COMBUSTION_TYPES:
        raise ValueError("연소구분을 선택하세요.")

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        INSERT INTO facility_energy
        (
            site,
            facility,
            energy_type,
            usage_unit,
            combustion_type,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(site, facility)
        DO UPDATE SET
            energy_type = excluded.energy_type,
            usage_unit = excluded.usage_unit,
            combustion_type = excluded.combustion_type,
            updated_at = excluded.updated_at
        """,
        (
            site,
            facility,
            energy_type,
            usage_unit,
            combustion_type,
            now_text,
            now_text,
        ),
    )

    conn.commit()
    conn.close()


def update_facility_energy_row(
    row_id,
    old_site,
    old_facility,
    new_facility,
    new_energy_type,
    new_usage_unit,
    new_combustion_type
):
    new_facility = new_facility.strip()

    if new_facility == "":
        raise ValueError("시설명을 입력하세요.")

    if new_usage_unit.strip() == "":
        raise ValueError("사용량 단위를 입력하세요.")

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        UPDATE facility_energy
        SET facility = ?,
            energy_type = ?,
            usage_unit = ?,
            combustion_type = ?,
            updated_at = ?
        WHERE id = ?
        """,
        (
            new_facility,
            new_energy_type,
            new_usage_unit,
            new_combustion_type,
            now_text,
            row_id,
        ),
    )

    cur.execute(
        """
        UPDATE activity_data
        SET facility = ?,
            energy_type = ?,
            usage_unit = ?,
            scope = ?,
            combustion_type = ?,
            updated_at = ?
        WHERE site = ?
          AND facility = ?
        """,
        (
            new_facility,
            new_energy_type,
            new_usage_unit,
            get_scope(new_energy_type),
            new_combustion_type,
            now_text,
            old_site,
            old_facility,
        ),
    )

    conn.commit()
    conn.close()

    recalculate_activity_data_for_site_facility(old_site, new_facility)


def delete_facility_energy_by_ids(ids):
    if len(ids) == 0:
        return 0

    placeholders = ",".join(["?"] * len(ids))

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        f"""
        DELETE FROM facility_energy
        WHERE id IN ({placeholders})
        """,
        ids,
    )

    deleted_count = cur.rowcount

    conn.commit()
    conn.close()

    return deleted_count


def load_activity_data():
    conn = get_connection()
    df = pd.read_sql_query(
        """
        SELECT
            id,
            year,
            month,
            site,
            facility,
            country,
            corporation,
            energy_type,
            usage_amount,
            usage_unit,
            scope,
            combustion_type,
            emission_amount,
            created_at,
            updated_at
        FROM activity_data
        ORDER BY year, month, site, facility
        """,
        conn,
    )
    conn.close()
    return df


def upsert_activity_data(
    year,
    month,
    site,
    facility,
    energy_type,
    usage_amount,
    usage_unit,
    combustion_type
):
    country = get_country(site)
    corporation = get_corporation(site)
    scope = get_scope(energy_type)

    emission_amount = calculate_emission(
        factor_year=year,
        site=site,
        energy_type=energy_type,
        usage_amount=usage_amount,
        combustion_type=combustion_type,
    )

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        INSERT INTO activity_data
        (
            year,
            month,
            site,
            facility,
            country,
            corporation,
            energy_type,
            usage_amount,
            usage_unit,
            scope,
            combustion_type,
            emission_amount,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(year, month, site, facility)
        DO UPDATE SET
            country = excluded.country,
            corporation = excluded.corporation,
            energy_type = excluded.energy_type,
            usage_amount = excluded.usage_amount,
            usage_unit = excluded.usage_unit,
            scope = excluded.scope,
            combustion_type = excluded.combustion_type,
            emission_amount = excluded.emission_amount,
            updated_at = excluded.updated_at
        """,
        (
            year,
            month,
            site,
            facility,
            country,
            corporation,
            energy_type,
            usage_amount,
            usage_unit,
            scope,
            combustion_type,
            emission_amount,
            now_text,
            now_text,
        ),
    )

    conn.commit()
    conn.close()

    return emission_amount


def delete_activity_data_by_ids(ids):
    if len(ids) == 0:
        return 0

    placeholders = ",".join(["?"] * len(ids))

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        f"""
        DELETE FROM activity_data
        WHERE id IN ({placeholders})
        """,
        ids,
    )

    deleted_count = cur.rowcount

    conn.commit()
    conn.close()

    return deleted_count

def recalculate_activity_data(year_filter=None):
    df = load_activity_data()

    if df.empty:
        return {
            "target_count": 0,
            "updated_count": 0,
            "missing_factor_count": 0,
        }

    if year_filter is not None:
        df = df[df["year"] == year_filter].copy()

    if df.empty:
        return {
            "target_count": 0,
            "updated_count": 0,
            "missing_factor_count": 0,
        }

    now_text = datetime.now().isoformat(timespec="seconds")

    target_count = len(df)
    updated_count = 0
    missing_factor_count = 0

    conn = get_connection()
    cur = conn.cursor()

    for _, row in df.iterrows():
        ncv, ef = get_factor(
            factor_year=int(row["year"]),
            country=str(row["country"]),
            energy_type=str(row["energy_type"]),
            combustion_type=str(row["combustion_type"]),
        )

        if ncv is None or ef is None:
            missing_factor_count += 1
            continue

        emission_amount = calculate_emission(
            factor_year=int(row["year"]),
            site=str(row["site"]),
            energy_type=str(row["energy_type"]),
            usage_amount=float(row["usage_amount"]),
            combustion_type=str(row["combustion_type"]),
        )

        cur.execute(
            """
            UPDATE activity_data
            SET emission_amount = ?,
                scope = ?,
                country = ?,
                corporation = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                emission_amount,
                get_scope(str(row["energy_type"])),
                get_country(str(row["site"])),
                get_corporation(str(row["site"])),
                now_text,
                int(row["id"]),
            ),
        )

        updated_count += 1

    conn.commit()
    conn.close()

    return {
        "target_count": target_count,
        "updated_count": updated_count,
        "missing_factor_count": missing_factor_count,
    }


def recalculate_activity_data_for_site_facility(site, facility):
    df = load_activity_data()

    target_df = df[
        (df["site"] == site)
        & (df["facility"] == facility)
    ].copy()

    for _, row in target_df.iterrows():
        emission_amount = calculate_emission(
            factor_year=int(row["year"]),
            site=row["site"],
            energy_type=row["energy_type"],
            usage_amount=row["usage_amount"],
            combustion_type=row["combustion_type"],
        )

        now_text = datetime.now().isoformat(timespec="seconds")

        conn = get_connection()
        cur = conn.cursor()

        cur.execute(
            """
            UPDATE activity_data
            SET emission_amount = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                emission_amount,
                now_text,
                int(row["id"]),
            ),
        )

        conn.commit()
        conn.close()


def update_activity_row(
    row_id,
    year,
    month,
    site,
    facility,
    energy_type,
    usage_amount,
    usage_unit,
    combustion_type
):
    if site not in SITES:
        raise ValueError("등록되지 않은 사업장입니다.")

    if energy_type not in ENERGY_TYPES:
        raise ValueError("등록되지 않은 에너지원입니다.")

    if combustion_type not in COMBUSTION_TYPES:
        raise ValueError("연소구분을 선택하세요.")

    add_facility_energy(
        site=site,
        facility=facility,
        energy_type=energy_type,
        usage_unit=usage_unit,
        combustion_type=combustion_type,
    )

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        DELETE FROM activity_data
        WHERE id = ?
        """,
        (row_id,),
    )

    conn.commit()
    conn.close()

    upsert_activity_data(
        year=year,
        month=month,
        site=site,
        facility=facility,
        energy_type=energy_type,
        usage_amount=usage_amount,
        usage_unit=usage_unit,
        combustion_type=combustion_type,
    )


def format_numeric_df(df):
    result_df = df.copy()

    numeric_columns = [
        "usage_amount",
        "emission_amount",
        "current_emission",
        "previous_emission",
        "change_amount",
        "change_rate",
    ]

    for col in numeric_columns:
        if col in result_df.columns:
            result_df[col] = result_df[col].round(3)

    return result_df


def render_facility_energy_page():
    st.subheader("시설-에너지원 관리")

    st.write(
        "시설 하나에는 하나의 에너지원만 연결됩니다. "
        "단위와 고정연소/이동연소 구분은 시설별로 관리합니다."
    )

    df = load_facility_energy()

    if df.empty:
        st.info("등록된 시설-에너지원 조합이 없습니다.")
    else:
        edit_df = df.copy()
        edit_df["삭제"] = False

        edit_df = edit_df[
            [
                "삭제",
                "id",
                "site",
                "facility",
                "energy_type",
                "usage_unit",
                "combustion_type",
                "created_at",
                "updated_at",
            ]
        ]

        edited_df = st.data_editor(
            edit_df,
            use_container_width=True,
            hide_index=True,
            num_rows="fixed",
            key="facility_energy_editor",
            column_order=[
                "삭제",
                "site",
                "facility",
                "energy_type",
                "usage_unit",
                "combustion_type",
                "created_at",
                "updated_at",
            ],
            disabled=[
                "id",
                "site",
                "created_at",
                "updated_at",
            ],
            column_config={
                "삭제": st.column_config.CheckboxColumn("삭제", default=False),
                "site": st.column_config.TextColumn("사업장"),
                "facility": st.column_config.TextColumn("시설명"),
                "energy_type": st.column_config.SelectboxColumn("에너지원", options=ENERGY_TYPES),
                "usage_unit": st.column_config.SelectboxColumn("사용량 단위", options=UNIT_OPTIONS),
                "combustion_type": st.column_config.SelectboxColumn("연소구분", options=COMBUSTION_TYPES),
                "created_at": st.column_config.TextColumn("생성일"),
                "updated_at": st.column_config.TextColumn("수정일"),
            }
        )

        if st.button("시설-에너지원 수정/삭제 저장", type="primary"):
            try:
                original_df = df.set_index("id")
                delete_ids = []

                for _, row in edited_df.iterrows():
                    row_id = int(row["id"])
                    should_delete = bool(row["삭제"])

                    old_site = str(original_df.loc[row_id, "site"])
                    old_facility = str(original_df.loc[row_id, "facility"])

                    new_facility = str(row["facility"]).strip()
                    new_energy_type = str(row["energy_type"])
                    new_usage_unit = str(row["usage_unit"])
                    new_combustion_type = str(row["combustion_type"])

                    if should_delete:
                        delete_ids.append(row_id)
                    else:
                        original_facility = str(original_df.loc[row_id, "facility"])
                        original_energy_type = str(original_df.loc[row_id, "energy_type"])
                        original_usage_unit = str(original_df.loc[row_id, "usage_unit"])
                        original_combustion_type = str(original_df.loc[row_id, "combustion_type"])

                        changed = (
                            original_facility != new_facility
                            or original_energy_type != new_energy_type
                            or original_usage_unit != new_usage_unit
                            or original_combustion_type != new_combustion_type
                        )

                        if changed:
                            update_facility_energy_row(
                                row_id=row_id,
                                old_site=old_site,
                                old_facility=old_facility,
                                new_facility=new_facility,
                                new_energy_type=new_energy_type,
                                new_usage_unit=new_usage_unit,
                                new_combustion_type=new_combustion_type,
                            )

                if delete_ids:
                    deleted_count = delete_facility_energy_by_ids(delete_ids)
                    st.success(f"{deleted_count}개 시설-에너지원 조합이 삭제되었습니다.")

                st.success("시설-에너지원 수정/삭제 작업이 완료되었습니다.")
                st.rerun()

            except sqlite3.IntegrityError:
                st.error(
                    "같은 사업장에 동일한 시설명이 이미 있습니다. "
                    "시설명은 같은 사업장 내에서 중복될 수 없습니다."
                )
            except Exception as e:
                st.error(f"시설-에너지원 수정/삭제 중 오류가 발생했습니다: {e}")

    st.divider()

    st.subheader("시설-에너지원 추가")

    with st.form("facility_energy_add_form"):
        site = st.selectbox("사업장", list(SITES.keys()))
        facility = st.text_input("시설명", value="")
        energy_type = st.selectbox("에너지원", ENERGY_TYPES)

        default_unit = get_default_unit(site, energy_type)
        usage_unit = st.selectbox(
            "사용량 단위",
            UNIT_OPTIONS,
            index=UNIT_OPTIONS.index(default_unit) if default_unit in UNIT_OPTIONS else 0
        )

        default_combustion_type = get_default_combustion_type(energy_type)
        combustion_type = st.selectbox(
            "연소구분",
            COMBUSTION_TYPES,
            index=COMBUSTION_TYPES.index(default_combustion_type)
        )

        scope = get_scope(energy_type)
        st.write(f"Scope 구분: {scope}")

        submitted = st.form_submit_button("시설-에너지원 추가")

        if submitted:
            try:
                add_facility_energy(
                    site=site,
                    facility=facility,
                    energy_type=energy_type,
                    usage_unit=usage_unit,
                    combustion_type=combustion_type,
                )
                st.success(
                    f"{site} / {facility} / {energy_type} / {usage_unit} / {combustion_type} 조합이 저장되었습니다."
                )
                st.rerun()
            except Exception as e:
                st.error(f"추가 중 오류가 발생했습니다: {e}")


def build_yearly_input_template(year, selected_site_filter):
    fe_df = load_facility_energy()

    if selected_site_filter != "전체":
        fe_df = fe_df[fe_df["site"] == selected_site_filter].copy()

    if fe_df.empty:
        return pd.DataFrame(
            columns=[
                "site",
                "facility",
                "energy_type",
                "usage_unit",
                "scope",
                "combustion_type",
                *MONTH_COLUMNS,
            ]
        )

    activity_df = load_activity_data()
    year_df = activity_df[activity_df["year"] == year].copy()

    rows = []

    for _, item in fe_df.iterrows():
        site = item["site"]
        facility = item["facility"]
        energy_type = item["energy_type"]
        usage_unit = item["usage_unit"]
        combustion_type = item["combustion_type"]
        scope = get_scope(energy_type)

        row_data = {
            "site": site,
            "facility": facility,
            "energy_type": energy_type,
            "usage_unit": usage_unit,
            "scope": scope,
            "combustion_type": combustion_type,
        }

        target_df = year_df[
            (year_df["site"] == site)
            & (year_df["facility"] == facility)
        ].copy()

        for month in range(1, 13):
            month_name = f"{month}월"

            month_values = target_df[
                target_df["month"] == month
            ]["usage_amount"]

            if not month_values.empty:
                row_data[month_name] = round(float(month_values.iloc[0]), 3)
            else:
                row_data[month_name] = 0.0

        rows.append(row_data)

    result_df = pd.DataFrame(rows)

    site_order = {site: index for index, site in enumerate(SITES.keys())}
    result_df["site_order"] = result_df["site"].map(site_order)
    result_df = result_df.sort_values(["site_order", "facility"]).drop(columns=["site_order"])

    return result_df


def save_yearly_input_data(year, edited_df):
    saved_count = 0

    for _, row in edited_df.iterrows():
        site = str(row["site"])
        facility = str(row["facility"]).strip()
        energy_type = str(row["energy_type"])
        usage_unit = str(row["usage_unit"])
        combustion_type = str(row["combustion_type"])

        if site not in SITES:
            continue

        if facility == "":
            continue

        if energy_type not in ENERGY_TYPES:
            continue

        add_facility_energy(
            site=site,
            facility=facility,
            energy_type=energy_type,
            usage_unit=usage_unit,
            combustion_type=combustion_type,
        )

        for month in range(1, 13):
            month_name = f"{month}월"
            usage_amount = round(float(row[month_name]), 3)

            upsert_activity_data(
                year=year,
                month=month,
                site=site,
                facility=facility,
                energy_type=energy_type,
                usage_amount=usage_amount,
                usage_unit=usage_unit,
                combustion_type=combustion_type,
            )

            saved_count += 1

    return saved_count
def make_yearly_excel_template(year, selected_site_filter):
    template_df = build_yearly_input_template(
        year=year,
        selected_site_filter=selected_site_filter,
    )

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        template_df.to_excel(
            writer,
            index=False,
            sheet_name="연간사용량"
        )

        guide_df = pd.DataFrame(
            {
                "항목": [
                    "입력 연도",
                    "작성 방법",
                    "주의 사항 1",
                    "주의 사항 2",
                    "주의 사항 3",
                    "주의 사항 4",
                ],
                "내용": [
                    year,
                    "연간사용량 시트에서 1월~12월 사용량만 입력한 뒤 업로드하세요.",
                    "site, facility, energy_type, usage_unit, scope, combustion_type 컬럼명은 변경하지 마세요.",
                    "사업장, 시설, 에너지원, 단위, 연소구분은 시설-에너지원 관리에 등록된 기준과 일치해야 합니다.",
                    "사용량은 0 이상의 숫자로 입력하세요. 소수점 셋째 자리까지 저장됩니다.",
                    "업로드 시 기존 동일 연도, 월, 사업장, 시설 데이터는 업데이트됩니다.",
                ],
            }
        )

        guide_df.to_excel(
            writer,
            index=False,
            sheet_name="작성가이드"
        )

        workbook = writer.book
        worksheet = writer.sheets["연간사용량"]

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 30)

        guide_sheet = writer.sheets["작성가이드"]

        for column_cells in guide_sheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))

            guide_sheet.column_dimensions[column_letter].width = min(max_length + 2, 80)

    output.seek(0)

    return output


def read_uploaded_yearly_excel(uploaded_file):
    try:
        uploaded_df = pd.read_excel(
            uploaded_file,
            sheet_name="연간사용량",
            engine="openpyxl"
        )
    except ValueError:
        uploaded_df = pd.read_excel(
            uploaded_file,
            engine="openpyxl"
        )

    uploaded_df.columns = [str(col).strip() for col in uploaded_df.columns]

    return uploaded_df


def validate_yearly_upload_df(uploaded_df):
    errors = []
    warnings = []

    required_columns = [
        "site",
        "facility",
        "energy_type",
        "usage_unit",
        "scope",
        "combustion_type",
        *MONTH_COLUMNS,
    ]

    missing_columns = [
        col for col in required_columns
        if col not in uploaded_df.columns
    ]

    if missing_columns:
        errors.append(
            "필수 컬럼이 누락되었습니다: "
            + ", ".join(missing_columns)
        )
        return errors, warnings

    if uploaded_df.empty:
        errors.append("업로드 파일에 데이터 행이 없습니다.")
        return errors, warnings

    fe_df = load_facility_energy()

    if fe_df.empty:
        errors.append("시설-에너지원 관리에 등록된 조합이 없습니다.")
        return errors, warnings

    registered_keys = set(
        zip(
            fe_df["site"].astype(str),
            fe_df["facility"].astype(str),
            fe_df["energy_type"].astype(str),
            fe_df["usage_unit"].astype(str),
            fe_df["combustion_type"].astype(str),
        )
    )

    duplicated_df = uploaded_df[
        uploaded_df.duplicated(
            subset=["site", "facility"],
            keep=False
        )
    ]

    if not duplicated_df.empty:
        duplicated_items = (
            duplicated_df["site"].astype(str)
            + " / "
            + duplicated_df["facility"].astype(str)
        ).drop_duplicates().tolist()

        errors.append(
            "업로드 파일 안에 같은 사업장/시설 조합이 중복되어 있습니다: "
            + ", ".join(duplicated_items[:10])
        )

    for index, row in uploaded_df.iterrows():
        excel_row_number = index + 2

        site = str(row["site"]).strip()
        facility = str(row["facility"]).strip()
        energy_type = str(row["energy_type"]).strip()
        usage_unit = str(row["usage_unit"]).strip()
        scope = str(row["scope"]).strip()
        combustion_type = str(row["combustion_type"]).strip()

        if site not in SITES:
            errors.append(
                f"{excel_row_number}행: 등록되지 않은 사업장입니다. site={site}"
            )

        if facility == "":
            errors.append(
                f"{excel_row_number}행: 시설명이 비어 있습니다."
            )

        if energy_type not in ENERGY_TYPES:
            errors.append(
                f"{excel_row_number}행: 등록되지 않은 에너지원입니다. energy_type={energy_type}"
            )

        if usage_unit not in UNIT_OPTIONS:
            errors.append(
                f"{excel_row_number}행: 등록되지 않은 사용량 단위입니다. usage_unit={usage_unit}"
            )

        if combustion_type not in COMBUSTION_TYPES:
            errors.append(
                f"{excel_row_number}행: 연소구분은 고정연소 또는 이동연소여야 합니다. combustion_type={combustion_type}"
            )

        if energy_type in ENERGY_TYPES:
            expected_scope = get_scope(energy_type)

            if scope != expected_scope:
                errors.append(
                    f"{excel_row_number}행: Scope가 에너지원 기준과 다릅니다. "
                    f"입력값={scope}, 기대값={expected_scope}"
                )

        key = (
            site,
            facility,
            energy_type,
            usage_unit,
            combustion_type,
        )

        if key not in registered_keys:
            errors.append(
                f"{excel_row_number}행: 시설-에너지원 관리에 등록된 조합과 일치하지 않습니다. "
                f"{site} / {facility} / {energy_type} / {usage_unit} / {combustion_type}"
            )

        for month_name in MONTH_COLUMNS:
            value = row[month_name]

            if pd.isna(value):
                warnings.append(
                    f"{excel_row_number}행 {month_name}: 빈 값은 0으로 처리됩니다."
                )
                continue

            try:
                numeric_value = float(value)
            except Exception:
                errors.append(
                    f"{excel_row_number}행 {month_name}: 숫자가 아닙니다. 입력값={value}"
                )
                continue

            if numeric_value < 0:
                errors.append(
                    f"{excel_row_number}행 {month_name}: 사용량은 0 이상이어야 합니다. 입력값={value}"
                )

    return errors, warnings


def normalize_yearly_upload_df(uploaded_df):
    result_df = uploaded_df.copy()

    text_columns = [
        "site",
        "facility",
        "energy_type",
        "usage_unit",
        "scope",
        "combustion_type",
    ]

    for col in text_columns:
        result_df[col] = result_df[col].astype(str).str.strip()

    for month_name in MONTH_COLUMNS:
        result_df[month_name] = (
            pd.to_numeric(result_df[month_name], errors="coerce")
            .fillna(0)
            .round(3)
        )

    return result_df




def render_yearly_input_page():
    st.subheader("연간 월별 사용량 입력")

    st.write(
        "연도 하나를 선택하고, 시설-에너지원별 1월부터 12월까지의 사용량을 입력합니다. "
        "엑셀 템플릿을 내려받아 작성한 뒤 업로드할 수도 있습니다."
    )

    col1, col2 = st.columns(2)

    with col1:
        selected_year = st.number_input(
            "입력 연도",
            min_value=START_YEAR,
            max_value=MAX_YEAR,
            value=get_default_year(),
            step=1
        )

    with col2:
        selected_site_filter = st.selectbox(
            "사업장 필터",
            ["전체"] + list(SITES.keys())
        )

    template_df = build_yearly_input_template(
        year=int(selected_year),
        selected_site_filter=selected_site_filter,
    )

    if template_df.empty:
        st.warning("등록된 시설-에너지원 조합이 없습니다. 먼저 시설-에너지원 관리 메뉴에서 조합을 등록하세요.")
        return

    st.divider()

    st.subheader("엑셀 템플릿 다운로드 / 업로드")

    template_file = make_yearly_excel_template(
        year=int(selected_year),
        selected_site_filter=selected_site_filter,
    )

    st.download_button(
        label="연간 월별 사용량 엑셀 템플릿 다운로드",
        data=template_file,
        file_name=f"GHG_연간월별사용량_템플릿_{int(selected_year)}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    uploaded_file = st.file_uploader(
        "작성한 연간 월별 사용량 엑셀 파일 업로드",
        type=["xlsx"]
    )

    if uploaded_file is not None:
        try:
            uploaded_df = read_uploaded_yearly_excel(uploaded_file)
            errors, warnings = validate_yearly_upload_df(uploaded_df)

            st.write("업로드 데이터 미리보기")
            preview_df = uploaded_df.copy()

            for month_name in MONTH_COLUMNS:
                if month_name in preview_df.columns:
                    preview_df[month_name] = pd.to_numeric(
                        preview_df[month_name],
                        errors="coerce"
                    ).fillna(0).round(3)

            st.dataframe(
                preview_df,
                use_container_width=True,
                hide_index=True
            )

            if warnings:
                with st.expander("업로드 경고 확인", expanded=False):
                    for warning in warnings[:100]:
                        st.warning(warning)

                    if len(warnings) > 100:
                        st.warning(f"경고가 {len(warnings)}건 있습니다. 상위 100건만 표시했습니다.")

            if errors:
                st.error("업로드 파일에 오류가 있어 저장할 수 없습니다.")

                with st.expander("업로드 오류 확인", expanded=True):
                    for error in errors[:100]:
                        st.error(error)

                    if len(errors) > 100:
                        st.error(f"오류가 {len(errors)}건 있습니다. 상위 100건만 표시했습니다.")
            else:
                normalized_df = normalize_yearly_upload_df(uploaded_df)

                total_upload_rows = len(normalized_df)
                total_activity_rows = total_upload_rows * 12

                st.success(
                    "업로드 파일 검증이 완료되었습니다. "
                    f"시설 행 수: {total_upload_rows}개, 저장 예정 월별 데이터 수: {total_activity_rows}건"
                )

                confirm_upload = st.checkbox(
                    "검증된 엑셀 데이터를 저장하겠습니다. 기존 동일 연도/월/사업장/시설 데이터는 업데이트됩니다."
                )

                if st.button("엑셀 업로드 데이터 저장", type="primary"):
                    if not confirm_upload:
                        st.warning("저장하려면 확인 체크박스를 선택하세요.")
                    else:
                        saved_count = save_yearly_input_data(
                            year=int(selected_year),
                            edited_df=normalized_df,
                        )

                        st.success(
                            f"{int(selected_year)}년 엑셀 업로드 데이터가 저장되었습니다. "
                            f"저장 또는 업데이트된 데이터 수: {saved_count}건"
                        )
                        st.rerun()

        except Exception as e:
            st.error(f"엑셀 업로드 처리 중 오류가 발생했습니다: {e}")

    st.divider()

    st.subheader("화면에서 직접 입력")

    month_column_config = {
        month_name: st.column_config.NumberColumn(
            month_name,
            min_value=0.0,
            step=0.001,
            format="%.3f"
        )
        for month_name in MONTH_COLUMNS
    }

    column_config = {
        "site": st.column_config.TextColumn("사업장"),
        "facility": st.column_config.TextColumn("시설"),
        "energy_type": st.column_config.TextColumn("에너지원"),
        "usage_unit": st.column_config.TextColumn("단위"),
        "scope": st.column_config.TextColumn("Scope"),
        "combustion_type": st.column_config.TextColumn("연소구분"),
    }
    column_config.update(month_column_config)

    edited_df = st.data_editor(
        template_df,
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        key="yearly_usage_editor",
        disabled=[
            "site",
            "facility",
            "energy_type",
            "usage_unit",
            "scope",
            "combustion_type",
        ],
        column_config=column_config
    )

    if st.button("연간 월별 사용량 저장", type="primary"):
        try:
            saved_count = save_yearly_input_data(
                year=int(selected_year),
                edited_df=edited_df,
            )

            st.success(
                f"{int(selected_year)}년 월별 사용량이 저장되었습니다. "
                f"저장 또는 업데이트된 데이터 수: {saved_count}건"
            )
            st.rerun()

        except Exception as e:
            st.error(f"저장 중 오류가 발생했습니다: {e}")



def render_factor_page():
    st.subheader("연도별·국가별·에너지원별·연소구분별 배출계수 관리")

    factors_df = load_factors()

    if factors_df.empty:
        st.info("등록된 배출계수가 없습니다.")
    else:
        filter_col1, filter_col2, filter_col3, filter_col4 = st.columns(4)

        with filter_col1:
            year_options = ["전체"] + sorted(factors_df["factor_year"].dropna().unique().tolist())
            selected_year = st.selectbox("연도 필터", year_options)

        with filter_col2:
            country_options = ["전체"] + sorted(factors_df["country"].dropna().unique().tolist())
            selected_country = st.selectbox("국가 필터", country_options)

        with filter_col3:
            energy_options = ["전체"] + ENERGY_TYPES
            selected_energy = st.selectbox("에너지원 필터", energy_options)

        with filter_col4:
            combustion_options = ["전체"] + COMBUSTION_TYPES
            selected_combustion = st.selectbox("연소구분 필터", combustion_options)

        view_df = factors_df.copy()

        if selected_year != "전체":
            view_df = view_df[view_df["factor_year"] == selected_year]

        if selected_country != "전체":
            view_df = view_df[view_df["country"] == selected_country]

        if selected_energy != "전체":
            view_df = view_df[view_df["energy_type"] == selected_energy]

        if selected_combustion != "전체":
            view_df = view_df[view_df["combustion_type"] == selected_combustion]

        display_factor_df = view_df[
            [
                "factor_year",
                "country",
                "energy_type",
                "combustion_type",
                "net_calorific_value",
                "emission_factor",
                "unit",
                "updated_at",
            ]
        ].copy()

        display_factor_df["net_calorific_value"] = display_factor_df["net_calorific_value"].round(6)
        display_factor_df["emission_factor"] = display_factor_df["emission_factor"].round(6)

        st.dataframe(
            display_factor_df,
            use_container_width=True,
            hide_index=True,
        )

    st.divider()

    st.subheader("배출계수 추가/수정")

    with st.form("factor_form"):
        factor_year = st.number_input(
            "계수 적용 연도",
            min_value=START_YEAR,
            max_value=MAX_YEAR,
            value=get_default_year(),
            step=1
        )

        country = st.selectbox(
            "국가",
            ["한국", "멕시코", "베트남", "중국"]
        )

        energy_type = st.selectbox(
            "에너지원",
            ENERGY_TYPES
        )

        if energy_type == "전기":
            combustion_type = st.selectbox(
                "연소구분",
                ["고정연소"]
            )
        else:
            combustion_type = st.selectbox(
                "연소구분",
                COMBUSTION_TYPES
            )

        ncv = st.number_input(
            "순발열량",
            min_value=0.0,
            value=0.0,
            step=0.000001,
            format="%.10f"
        )

        ef = st.number_input(
            "배출계수",
            min_value=0.0,
            value=0.0,
            step=0.000001,
            format="%.10f"
        )

        unit = st.text_input(
            "계수 단위 설명",
            value=""
        )

        submitted = st.form_submit_button("계수 저장")

        if submitted:
            update_factor(
                factor_year=int(factor_year),
                country=country,
                energy_type=energy_type,
                combustion_type=combustion_type,
                ncv=ncv,
                ef=ef,
                unit=unit,
            )
            st.success("배출계수가 저장되었습니다.")

    st.divider()

    st.subheader("배출량 재계산")

    st.write(
        "배출계수를 수정한 뒤 이미 저장된 사용량 데이터의 배출량을 다시 계산합니다. "
        "재계산 전에는 DB 백업을 먼저 권장합니다."
    )

    st.warning(
        "재계산을 실행하면 activity_data 테이블의 emission_amount 값이 현재 배출계수 기준으로 업데이트됩니다."
    )

    recalc_col1, recalc_col2 = st.columns(2)

    with recalc_col1:
        recalc_year = st.number_input(
            "재계산 연도",
            min_value=START_YEAR,
            max_value=MAX_YEAR,
            value=get_default_year(),
            step=1,
            key="factor_recalc_year"
        )

        confirm_year_recalc = st.checkbox(
            f"{int(recalc_year)}년 데이터를 재계산합니다.",
            key="confirm_year_recalc"
        )

        if st.button(
            "선택 연도 재계산",
            type="primary",
            disabled=not confirm_year_recalc
        ):
            result = recalculate_activity_data(
                year_filter=int(recalc_year)
            )

            st.success(
                f"{int(recalc_year)}년 재계산이 완료되었습니다. "
                f"대상: {result['target_count']}건, "
                f"업데이트: {result['updated_count']}건, "
                f"배출계수 누락: {result['missing_factor_count']}건"
            )

    with recalc_col2:
        confirm_all_recalc = st.checkbox(
            "전체 연도 데이터를 재계산합니다.",
            key="confirm_all_recalc"
        )

        if st.button(
            "전체 데이터 재계산",
            disabled=not confirm_all_recalc
        ):
            result = recalculate_activity_data(
                year_filter=None
            )

            st.success(
                "전체 데이터 재계산이 완료되었습니다. "
                f"대상: {result['target_count']}건, "
                f"업데이트: {result['updated_count']}건, "
                f"배출계수 누락: {result['missing_factor_count']}건"
            )
        


def render_data_edit_page():
    st.subheader("입력 데이터 조회")

    st.write(
        "저장된 월별 사용량과 배출량 데이터를 조회합니다. "
        "사용량 수정은 '연간 월별 사용량 입력' 메뉴에서 진행하고, "
        "시설명/에너지원/단위/연소구분 수정은 '시설-에너지원 관리' 메뉴에서 진행하세요."
    )

    df = load_activity_data()

    if df.empty:
        st.info("입력된 데이터가 없습니다.")
        return

    filter_row1_col1, filter_row1_col2, filter_row1_col3 = st.columns(3)

    with filter_row1_col1:
        year_options = ["전체"] + sorted(
            df["year"].dropna().astype(int).unique().tolist()
        )

        selected_year = st.selectbox(
            "연도 필터",
            year_options,
            key="data_view_year_filter"
        )

    filtered_for_month = df.copy()

    if selected_year != "전체":
        filtered_for_month = filtered_for_month[
            filtered_for_month["year"] == selected_year
        ].copy()

    with filter_row1_col2:
        available_months = sorted(
            filtered_for_month["month"].dropna().astype(int).unique().tolist()
        )

        month_options = ["전체"] + available_months

        selected_month = st.selectbox(
            "월 필터",
            month_options,
            key="data_view_month_filter"
        )

    filtered_for_site = filtered_for_month.copy()

    if selected_month != "전체":
        filtered_for_site = filtered_for_site[
            filtered_for_site["month"] == selected_month
        ].copy()

    with filter_row1_col3:
        available_sites = [
            site
            for site in SITES.keys()
            if site in filtered_for_site["site"].dropna().astype(str).unique().tolist()
        ]

        site_options = ["전체"] + available_sites

        selected_site = st.selectbox(
            "사업장 필터",
            site_options,
            key="data_view_site_filter"
        )

    filter_row2_col1, filter_row2_col2, filter_row2_col3 = st.columns(3)

    filtered_for_facility = filtered_for_site.copy()

    if selected_site != "전체":
        filtered_for_facility = filtered_for_facility[
            filtered_for_facility["site"] == selected_site
        ].copy()

    with filter_row2_col1:
        available_facilities = sorted(
            filtered_for_facility["facility"].dropna().astype(str).unique().tolist()
        )

        facility_options = ["전체"] + available_facilities

        selected_facility = st.selectbox(
            "시설 필터",
            facility_options,
            key="data_view_facility_filter"
        )

    filtered_for_energy = filtered_for_facility.copy()

    if selected_facility != "전체":
        filtered_for_energy = filtered_for_energy[
            filtered_for_energy["facility"] == selected_facility
        ].copy()

    with filter_row2_col2:
        available_energies = [
            energy_type
            for energy_type in ENERGY_TYPES
            if energy_type in filtered_for_energy["energy_type"].dropna().astype(str).unique().tolist()
        ]

        energy_options = ["전체"] + available_energies

        selected_energy = st.selectbox(
            "에너지원 필터",
            energy_options,
            key="data_view_energy_filter"
        )

    filtered_for_combustion = filtered_for_energy.copy()

    if selected_energy != "전체":
        filtered_for_combustion = filtered_for_combustion[
            filtered_for_combustion["energy_type"] == selected_energy
        ].copy()

    with filter_row2_col3:
        available_combustion_types = [
            combustion_type
            for combustion_type in COMBUSTION_TYPES
            if combustion_type in filtered_for_combustion["combustion_type"].dropna().astype(str).unique().tolist()
        ]

        combustion_options = ["전체"] + available_combustion_types

        selected_combustion = st.selectbox(
            "연소구분 필터",
            combustion_options,
            key="data_view_combustion_filter"
        )

    filter_row3_col1, filter_row3_col2, filter_row3_col3 = st.columns(3)

    filtered_for_scope = filtered_for_combustion.copy()

    if selected_combustion != "전체":
        filtered_for_scope = filtered_for_scope[
            filtered_for_scope["combustion_type"] == selected_combustion
        ].copy()

    with filter_row3_col1:
        available_scopes = sorted(
            filtered_for_scope["scope"].dropna().astype(str).unique().tolist()
        )

        scope_options = ["전체"] + available_scopes

        selected_scope = st.selectbox(
            "Scope 필터",
            scope_options,
            key="data_view_scope_filter"
        )

    view_df = filtered_for_scope.copy()

    if selected_scope != "전체":
        view_df = view_df[
            view_df["scope"] == selected_scope
        ].copy()

    if view_df.empty:
        st.warning("필터 조건에 해당하는 데이터가 없습니다.")
        return

    total_usage_amount = view_df["usage_amount"].sum()
    total_emission_amount = view_df["emission_amount"].sum()
    unique_facility_count = len(
        view_df[
            [
                "site",
                "facility",
            ]
        ].drop_duplicates()
    )

    summary_col1, summary_col2, summary_col3, summary_col4 = st.columns(4)

    with summary_col1:
        st.metric("조회 데이터 수", f"{len(view_df):,}건")

    with summary_col2:
        st.metric("조회 시설 수", f"{unique_facility_count:,}개")

    with summary_col3:
        st.metric("조회 사용량 합계", f"{total_usage_amount:,.3f}")

    with summary_col4:
        st.metric("조회 배출량 합계", f"{total_emission_amount:,.3f} tCO₂e")

        display_df = view_df[
        [
            "year",
            "month",
            "site",
            "facility",
            "country",
            "energy_type",
            "usage_amount",
            "usage_unit",
            "scope",
            "combustion_type",
            "emission_amount",
            "updated_at",
            "id",
        ]
    ].copy()

        display_df = display_df.rename(
        columns={
            "year": "연도",
            "month": "월",
            "site": "사업장",
            "facility": "시설",
            "country": "국가",
            "energy_type": "에너지원",
            "usage_amount": "사용량",
            "usage_unit": "단위",
            "scope": "Scope",
            "combustion_type": "연소구분",
            "emission_amount": "배출량(tCO₂e)",
            "updated_at": "수정일",
            "id": "ID",
        }
    )

    display_df["사용량"] = display_df["사용량"].round(3)
    display_df["배출량(tCO₂e)"] = display_df["배출량(tCO₂e)"].round(3)

    display_df = display_df.sort_values(
        [
            "연도",
            "월",
            "사업장",
            "시설",
            "에너지원",
        ]
    )

    st.divider()

    st.subheader("조회 결과")

    st.caption(
        "ID는 관리용 식별번호입니다. "
    )

    st.dataframe(
        display_df,
        use_container_width=True,
        hide_index=True,
    )

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        display_df.to_excel(
            writer,
            index=False,
            sheet_name="입력데이터조회"
        )

        worksheet = writer.sheets["입력데이터조회"]

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 35)

    output.seek(0)

    st.download_button(
        label="조회 결과 엑셀 다운로드",
        data=output,
        file_name=f"GHG_입력데이터조회_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def get_monthly_emission_values(df, year, site=None):
    if df.empty:
        return [0.0] * 12

    target_df = df[df["year"] == year].copy()

    if site is not None and site != "Total":
        target_df = target_df[target_df["site"] == site].copy()

    month_values = []

    for month in range(1, 13):
        month_sum = target_df[
            target_df["month"] == month
        ]["emission_amount"].sum()

        month_values.append(round(float(month_sum), 3))

    return month_values


def calculate_change_rate_values(previous_values, current_values):
    rate_values = []

    for previous_value, current_value in zip(previous_values, current_values):
        if previous_value == 0:
            rate_values.append("")
        else:
            change_rate = (current_value - previous_value) / previous_value * 100
            rate_values.append(round(float(change_rate), 1))

    return rate_values


def build_monthly_emission_comparison_table(base_year):
    df = load_activity_data()

    previous_year = base_year - 1
    current_year = base_year

    rows = []

    group_items = ["Total"] + list(SITES.keys())

    for group_name in group_items:
        title_row = {
            "row_type": "title",
            "사업부": group_name,
            "구분": "",
            "배출량(tCO₂e)": "",
        }

        for month_name in MONTH_COLUMNS:
            title_row[month_name] = ""

        title_row["감축목표"] = ""
        rows.append(title_row)

        previous_values = get_monthly_emission_values(
            df=df,
            year=previous_year,
            site=group_name,
        )

        current_values = get_monthly_emission_values(
            df=df,
            year=current_year,
            site=group_name,
        )

        rate_values = calculate_change_rate_values(
            previous_values=previous_values,
            current_values=current_values,
        )

        previous_total = round(sum(previous_values), 3)
        current_total = round(sum(current_values), 3)

        if previous_total == 0:
            total_rate = ""
        else:
            total_rate = round((current_total - previous_total) / previous_total * 100, 1)

        previous_row = {
            "row_type": "previous",
            "사업부": "",
            "구분": str(previous_year),
            "배출량(tCO₂e)": previous_total,
        }

        current_row = {
            "row_type": "current",
            "사업부": "",
            "구분": str(current_year),
            "배출량(tCO₂e)": current_total,
        }

        rate_row = {
            "row_type": "rate",
            "사업부": "",
            "구분": "증감률(%)",
            "배출량(tCO₂e)": total_rate,
        }

        for index, month_name in enumerate(MONTH_COLUMNS):
            previous_row[month_name] = previous_values[index]
            current_row[month_name] = current_values[index]
            rate_row[month_name] = rate_values[index]

        previous_row["감축목표"] = ""
        current_row["감축목표"] = ""
        rate_row["감축목표"] = ""

        rows.append(previous_row)
        rows.append(current_row)
        rows.append(rate_row)

        blank_row = {
            "row_type": "blank",
            "사업부": "",
            "구분": "",
            "배출량(tCO₂e)": "",
        }

        for month_name in MONTH_COLUMNS:
            blank_row[month_name] = ""

        blank_row["감축목표"] = ""
        rows.append(blank_row)

    result_df = pd.DataFrame(rows)

    display_columns = [
        "사업부",
        "구분",
        "배출량(tCO₂e)",
        *MONTH_COLUMNS,
        "감축목표",
        "row_type",
    ]

    result_df = result_df[display_columns]

    return result_df


def format_monthly_db_cell(value, is_rate_row=False):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return str(value)

    if numeric_value == 0:
        if is_rate_row:
            return "0.0"
        return "-"

    if is_rate_row:
        return f"{numeric_value:,.1f}"

    return f"{numeric_value:,.3f}"


def get_rate_value_class(value):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return ""

    if numeric_value > 0:
        return " rate-increase"

    if numeric_value < 0:
        return " rate-decrease"

    return " rate-zero"


def build_monthly_emission_html_table(df):
    table_columns = [
        "사업부",
        "구분",
        "배출량(tCO₂e)",
        *MONTH_COLUMNS,
        "감축목표",
    ]

    html = """
    <style>
        .monthly-db-table-wrapper {
            width: 100%;
            overflow-x: auto;
            margin-top: 12px;
            margin-bottom: 24px;
        }

        .monthly-db-table {
            width: 100%;
            border-collapse: collapse;
            font-size: 13px;
            table-layout: fixed;
        }

        .monthly-db-table th {
            background-color: #F3F4F6;
            color: #111827;
            border: 1px solid #D1D5DB;
            padding: 6px 6px;
            text-align: center;
            font-weight: 700;
            white-space: nowrap;
        }

        .monthly-db-table td {
            border: 1px solid #D1D5DB;
            padding: 5px 6px;
            text-align: right;
            white-space: nowrap;
        }

        .monthly-db-table td.col-business {
            text-align: left;
            font-weight: 600;
            width: 90px;
        }

        .monthly-db-table td.col-type {
            text-align: center;
            width: 80px;
        }

        .monthly-db-table th.col-business {
            width: 90px;
        }

        .monthly-db-table th.col-type {
            width: 80px;
        }

        .monthly-db-table th.col-total {
            width: 110px;
        }

        .monthly-db-table th.col-month {
            width: 75px;
        }

        .monthly-db-table th.col-target {
            width: 80px;
        }

        .monthly-db-table tr.title-row td {
            background-color: #E5E7EB;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .monthly-db-table tr.total-title-row td {
            background-color: #D9EAF7;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .monthly-db-table tr.rate-row td {
            background-color: #FFF2CC;
            font-weight: 700;
        }

        .monthly-db-table tr.blank-row td {
            background-color: #FFFFFF;
            height: 12px;
            border-left-color: #FFFFFF;
            border-right-color: #FFFFFF;
        }

        .monthly-db-table td.rate-increase {
            color: #B91C1C;
            font-weight: 800;
        }

        .monthly-db-table td.rate-decrease {
            color: #1D4ED8;
            font-weight: 800;
        }

        .monthly-db-table td.rate-zero {
            color: #111827;
            font-weight: 700;
        }
    </style>
    """

    html += '<div class="monthly-db-table-wrapper">'
    html += '<table class="monthly-db-table">'
    html += "<thead><tr>"

    for col in table_columns:
        class_name = ""

        if col == "사업부":
            class_name = "col-business"
        elif col == "구분":
            class_name = "col-type"
        elif col == "배출량(tCO₂e)":
            class_name = "col-total"
        elif col in MONTH_COLUMNS:
            class_name = "col-month"
        elif col == "감축목표":
            class_name = "col-target"

        html += f'<th class="{class_name}">{col}</th>'

    html += "</tr></thead>"
    html += "<tbody>"

    for _, row in df.iterrows():
        row_type = str(row["row_type"])
        business = str(row["사업부"])

        row_classes = []

        if row_type == "title":
            row_classes.append("title-row")

            if business == "Total":
                row_classes.append("total-title-row")

        if row_type == "rate":
            row_classes.append("rate-row")

        if row_type == "blank":
            row_classes.append("blank-row")

        row_class_text = " ".join(row_classes)

        html += f'<tr class="{row_class_text}">'

        if row_type == "title":
            html += f'<td class="col-business" colspan="{len(table_columns)}">{business}</td>'
            html += "</tr>"
            continue

        for col in table_columns:
            value = row[col]
            is_rate_row = row_type == "rate"
            rate_class = get_rate_value_class(value) if is_rate_row and col not in ["사업부", "구분", "감축목표"] else ""

            if col == "사업부":
                html += f'<td class="col-business">{value}</td>'
            elif col == "구분":
                html += f'<td class="col-type">{value}</td>'
            elif col == "감축목표":
                html += f"<td>{value}</td>"
            else:
                html += f'<td class="{rate_class.strip()}">{format_monthly_db_cell(value, is_rate_row=is_rate_row)}</td>'

        html += "</tr>"

    html += "</tbody></table></div>"

    return html


def make_monthly_emission_excel_file(df, base_year):
    output = BytesIO()

    export_df = df.drop(columns=["row_type"]).copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(
            writer,
            index=False,
            sheet_name="월별배출량DB"
        )

        workbook = writer.book
        worksheet = writer.sheets["월별배출량DB"]

        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="F3F4F6")
        title_fill = PatternFill("solid", fgColor="E5E7EB")
        total_title_fill = PatternFill("solid", fgColor="D9EAF7")
        rate_fill = PatternFill("solid", fgColor="FFF2CC")
        white_fill = PatternFill("solid", fgColor="FFFFFF")

        border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        red_font = Font(color="B91C1C", bold=True)
        blue_font = Font(color="1D4ED8", bold=True)
        black_bold_font = Font(color="111827", bold=True)
        header_font = Font(color="111827", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        visible_columns = export_df.columns.tolist()

        total_col_index = visible_columns.index("배출량(tCO₂e)") + 1
        month_col_indexes = [
            visible_columns.index(month_name) + 1
            for month_name in MONTH_COLUMNS
        ]

        for excel_row_index, row in enumerate(df.itertuples(index=False), start=2):
            row_dict = row._asdict()
            row_type = row_dict["row_type"]
            business = row_dict["사업부"]

            if row_type == "title":
                worksheet.merge_cells(
                    start_row=excel_row_index,
                    start_column=1,
                    end_row=excel_row_index,
                    end_column=len(visible_columns)
                )

            for col_index in range(1, len(visible_columns) + 1):
                cell = worksheet.cell(row=excel_row_index, column=col_index)
                cell.border = border
                cell.alignment = Alignment(horizontal="right", vertical="center")

                if col_index in [1, 2]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                if row_type == "title":
                    if business == "Total":
                        cell.fill = total_title_fill
                    else:
                        cell.fill = title_fill

                    cell.font = black_bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif row_type == "rate":
                    cell.fill = rate_fill
                    cell.font = black_bold_font

                    if col_index in [total_col_index, *month_col_indexes]:
                        value = cell.value

                        if isinstance(value, (int, float)):
                            if value > 0:
                                cell.font = red_font
                            elif value < 0:
                                cell.font = blue_font
                            else:
                                cell.font = black_bold_font

                elif row_type == "blank":
                    cell.fill = white_fill

                if col_index in [total_col_index, *month_col_indexes]:
                    if isinstance(cell.value, (int, float)):
                        if row_type == "rate":
                            cell.number_format = '#,##0.0'
                        else:
                            cell.number_format = '#,##0.000'

            if row_type == "blank":
                worksheet.row_dimensions[excel_row_index].height = 8

        column_widths = {
            "A": 14,
            "B": 14,
            "C": 16,
        }

        for col_letter, width in column_widths.items():
            worksheet.column_dimensions[col_letter].width = width

        for col_index in range(4, 16):
            col_letter = get_column_letter(col_index)
            worksheet.column_dimensions[col_letter].width = 12

        worksheet.column_dimensions["P"].width = 12

        worksheet.freeze_panes = "D2"

    output.seek(0)

    return output


def render_monthly_emission_db_page():
    st.subheader("온실가스 배출량 월별 DB")

    st.write(
        "전년도와 당해연도의 월별 배출량을 Total 및 사업장별로 비교합니다. "
        "전년도 배출량이 0인 월의 증감률은 빈칸으로 표시합니다."
    )

    selected_year = st.number_input(
        "기준연도",
        min_value=START_YEAR + 1,
        max_value=MAX_YEAR,
        value=max(get_default_year(), START_YEAR + 1),
        step=1
    )

    previous_year = int(selected_year) - 1

    st.caption(
        f"비교 기준: {previous_year}년 대비 {int(selected_year)}년"
    )

    comparison_df = build_monthly_emission_comparison_table(
        base_year=int(selected_year)
    )

    if comparison_df.empty:
        st.info("표시할 데이터가 없습니다.")
        return

    html_table = build_monthly_emission_html_table(comparison_df)

    st.markdown(
        html_table,
        unsafe_allow_html=True
    )

    excel_file = make_monthly_emission_excel_file(
        df=comparison_df,
        base_year=int(selected_year)
    )

    st.download_button(
        label="온실가스 배출량 월별 DB 엑셀 다운로드",
        data=excel_file,
        file_name=f"GHG_월별배출량_DB_{int(selected_year)}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def get_site_monthly_total_emission_values(df, site, year):
    if df.empty:
        return [0.0] * 12

    target_df = df[
        (df["site"] == site)
        & (df["year"] == year)
    ].copy()

    month_values = []

    for month in range(1, 13):
        month_sum = target_df[
            target_df["month"] == month
        ]["emission_amount"].sum()

        month_values.append(round(float(month_sum), 3))

    return month_values


def get_facility_monthly_values(df, site, facility, energy_type, usage_unit, year):
    if df.empty:
        return {
            "usage_values": [0.0] * 12,
            "emission_values": [0.0] * 12,
        }

    target_df = df[
        (df["site"] == site)
        & (df["facility"] == facility)
        & (df["energy_type"] == energy_type)
        & (df["usage_unit"] == usage_unit)
        & (df["year"] == year)
    ].copy()

    usage_values = []
    emission_values = []

    for month in range(1, 13):
        month_df = target_df[target_df["month"] == month].copy()

        usage_sum = month_df["usage_amount"].sum()
        emission_sum = month_df["emission_amount"].sum()

        usage_values.append(round(float(usage_sum), 3))
        emission_values.append(round(float(emission_sum), 3))

    return {
        "usage_values": usage_values,
        "emission_values": emission_values,
    }


def build_site_db_table(site, base_year):
    df = load_activity_data()

    year_1 = base_year - 2
    year_2 = base_year - 1
    year_3 = base_year

    years = [year_1, year_2, year_3]

    rows = []

    title_row = {
        "row_type": "section_title",
        "시설": "온실가스 배출량",
        "구분": "",
    }

    for month_name in MONTH_COLUMNS:
        title_row[month_name] = ""

    title_row["합계"] = ""
    rows.append(title_row)

    for year in years:
        emission_values = get_site_monthly_total_emission_values(
            df=df,
            site=site,
            year=year,
        )

        row = {
            "row_type": "total_emission",
            "시설": f"{str(year)[-2:]}년 실적",
            "구분": "배출량(tCO₂e)",
        }

        for index, month_name in enumerate(MONTH_COLUMNS):
            row[month_name] = emission_values[index]

        row["합계"] = round(sum(emission_values), 3)
        rows.append(row)

    previous_values = get_site_monthly_total_emission_values(
        df=df,
        site=site,
        year=year_2,
    )

    current_values = get_site_monthly_total_emission_values(
        df=df,
        site=site,
        year=year_3,
    )

    diff_values = [
        round(current - previous, 3)
        for previous, current in zip(previous_values, current_values)
    ]

    diff_row = {
        "row_type": "diff",
        "시설": "실적 비교",
        "구분": "전년대비 증감(tCO₂e)",
    }

    for index, month_name in enumerate(MONTH_COLUMNS):
        diff_row[month_name] = diff_values[index]

    diff_row["합계"] = round(sum(current_values) - sum(previous_values), 3)
    rows.append(diff_row)

    blank_row = {
        "row_type": "blank",
        "시설": "",
        "구분": "",
    }

    for month_name in MONTH_COLUMNS:
        blank_row[month_name] = ""

    blank_row["합계"] = ""
    rows.append(blank_row)

    site_df = df[df["site"] == site].copy()

    if not site_df.empty:
        facility_items = (
            site_df[
                [
                    "facility",
                    "energy_type",
                    "usage_unit",
                    "combustion_type",
                ]
            ]
            .drop_duplicates()
            .sort_values(["facility", "energy_type"])
        )

        for _, item in facility_items.iterrows():
            facility = str(item["facility"])
            energy_type = str(item["energy_type"])
            usage_unit = str(item["usage_unit"])

            facility_title = {
                "row_type": "facility_title",
                "시설": f"{facility} / {energy_type}",
                "구분": "",
            }

            for month_name in MONTH_COLUMNS:
                facility_title[month_name] = ""

            facility_title["합계"] = ""
            rows.append(facility_title)

            year_emission_map = {}

            for year in years:
                values = get_facility_monthly_values(
                    df=df,
                    site=site,
                    facility=facility,
                    energy_type=energy_type,
                    usage_unit=usage_unit,
                    year=year,
                )

                usage_values = values["usage_values"]
                emission_values = values["emission_values"]

                year_emission_map[year] = emission_values

                usage_row = {
                    "row_type": "usage",
                    "시설": f"{str(year)[-2:]}년 실적",
                    "구분": f"사용량({usage_unit})",
                }

                for index, month_name in enumerate(MONTH_COLUMNS):
                    usage_row[month_name] = usage_values[index]

                usage_row["합계"] = round(sum(usage_values), 3)
                rows.append(usage_row)

                emission_row = {
                    "row_type": "emission",
                    "시설": "",
                    "구분": "배출량(tCO₂e)",
                }

                for index, month_name in enumerate(MONTH_COLUMNS):
                    emission_row[month_name] = emission_values[index]

                emission_row["합계"] = round(sum(emission_values), 3)
                rows.append(emission_row)

            previous_facility_values = year_emission_map[year_2]
            current_facility_values = year_emission_map[year_3]

            facility_diff_values = [
                round(current - previous, 3)
                for previous, current in zip(previous_facility_values, current_facility_values)
            ]

            facility_diff_row = {
                "row_type": "diff",
                "시설": "실적 비교",
                "구분": "전년대비 증감(tCO₂e)",
            }

            for index, month_name in enumerate(MONTH_COLUMNS):
                facility_diff_row[month_name] = facility_diff_values[index]

            facility_diff_row["합계"] = round(
                sum(current_facility_values) - sum(previous_facility_values),
                3
            )

            rows.append(facility_diff_row)

            facility_blank_row = {
                "row_type": "blank",
                "시설": "",
                "구분": "",
            }

            for month_name in MONTH_COLUMNS:
                facility_blank_row[month_name] = ""

            facility_blank_row["합계"] = ""
            rows.append(facility_blank_row)

    result_df = pd.DataFrame(rows)

    display_columns = [
        "시설",
        "구분",
        *MONTH_COLUMNS,
        "합계",
        "row_type",
    ]

    result_df = result_df[display_columns]

    return result_df



def format_site_db_cell(value):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return str(value)

    if numeric_value == 0:
        return "-"

    return f"{numeric_value:,.3f}"


def get_diff_value_class(value):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return ""

    if numeric_value > 0:
        return " diff-increase"

    if numeric_value < 0:
        return " diff-decrease"

    return " diff-zero"


def build_site_db_html_table(df):
    table_columns = [
        "시설",
        "구분",
        *MONTH_COLUMNS,
        "합계",
    ]

    html = """
    <style>
        .site-db-table-wrapper {
            width: 100%;
            overflow-x: auto;
            margin-top: 12px;
            margin-bottom: 24px;
        }

        .site-db-table {
            width: 100%;
            border-collapse: collapse;
            font-size: 12px;
            table-layout: fixed;
        }

        .site-db-table th {
            background-color: #F3F4F6;
            color: #111827;
            border: 1px solid #D1D5DB;
            padding: 6px 5px;
            text-align: center;
            font-weight: 700;
            white-space: nowrap;
        }

        .site-db-table td {
            border: 1px solid #D1D5DB;
            padding: 5px 5px;
            text-align: right;
            white-space: nowrap;
        }

        .site-db-table th.col-facility,
        .site-db-table td.col-facility {
            width: 180px;
            text-align: left;
        }

        .site-db-table th.col-type,
        .site-db-table td.col-type {
            width: 130px;
            text-align: center;
        }

        .site-db-table th.col-month {
            width: 76px;
        }

        .site-db-table th.col-total {
            width: 95px;
        }

        .site-db-table tr.section-title-row td {
            background-color: #D9EAF7;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .site-db-table tr.facility-title-row td {
            background-color: #E5E7EB;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .site-db-table tr.diff-row td {
            background-color: #FFF2CC;
            font-weight: 700;
        }

        .site-db-table tr.blank-row td {
            background-color: #FFFFFF;
            height: 10px;
            border-left-color: #FFFFFF;
            border-right-color: #FFFFFF;
        }

        .site-db-table td.diff-increase {
            color: #B91C1C;
            font-weight: 800;
        }

        .site-db-table td.diff-decrease {
            color: #1D4ED8;
            font-weight: 800;
        }

        .site-db-table td.diff-zero {
            color: #111827;
            font-weight: 700;
        }
    </style>
    """

    html += '<div class="site-db-table-wrapper">'
    html += '<table class="site-db-table">'
    html += "<thead><tr>"

    for col in table_columns:
        class_name = ""

        if col == "시설":
            class_name = "col-facility"
        elif col == "구분":
            class_name = "col-type"
        elif col in MONTH_COLUMNS:
            class_name = "col-month"
        elif col == "합계":
            class_name = "col-total"

        html += f'<th class="{class_name}">{col}</th>'

    html += "</tr></thead>"
    html += "<tbody>"

    for _, row in df.iterrows():
        row_type = str(row["row_type"])

        row_classes = []

        if row_type == "section_title":
            row_classes.append("section-title-row")

        if row_type == "facility_title":
            row_classes.append("facility-title-row")

        if row_type == "diff":
            row_classes.append("diff-row")

        if row_type == "blank":
            row_classes.append("blank-row")

        row_class_text = " ".join(row_classes)

        html += f'<tr class="{row_class_text}">'

        if row_type in ["section_title", "facility_title"]:
            title_text = str(row["시설"])

            html += f'<td class="col-facility" colspan="{len(table_columns)}">{title_text}</td>'
            html += "</tr>"
            continue

        for col in table_columns:
            value = row[col]
            is_diff_row = row_type == "diff"
            diff_class = get_diff_value_class(value) if is_diff_row and col not in ["시설", "구분"] else ""

            if col == "시설":
                html += f'<td class="col-facility">{value}</td>'
            elif col == "구분":
                html += f'<td class="col-type">{value}</td>'
            else:
                html += f'<td class="{diff_class.strip()}">{format_site_db_cell(value)}</td>'

        html += "</tr>"

    html += "</tbody></table></div>"

    return html



def make_site_db_excel_file(df, site, base_year):
    output = BytesIO()

    export_df = df.drop(columns=["row_type"]).copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(
            writer,
            index=False,
            sheet_name=f"{site}_사업장DB"
        )

        worksheet = writer.sheets[f"{site}_사업장DB"]

        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="F3F4F6")
        section_fill = PatternFill("solid", fgColor="D9EAF7")
        facility_fill = PatternFill("solid", fgColor="E5E7EB")
        diff_fill = PatternFill("solid", fgColor="FFF2CC")
        white_fill = PatternFill("solid", fgColor="FFFFFF")

        border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        red_font = Font(color="B91C1C", bold=True)
        blue_font = Font(color="1D4ED8", bold=True)
        black_bold_font = Font(color="111827", bold=True)
        header_font = Font(color="111827", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        visible_columns = export_df.columns.tolist()
        month_col_indexes = [
            visible_columns.index(month_name) + 1
            for month_name in MONTH_COLUMNS
        ]
        total_col_index = visible_columns.index("합계") + 1

        for excel_row_index, row in enumerate(df.itertuples(index=False), start=2):
            row_dict = row._asdict()
            row_type = row_dict["row_type"]

            if row_type in ["section_title", "facility_title"]:
                worksheet.merge_cells(
                    start_row=excel_row_index,
                    start_column=1,
                    end_row=excel_row_index,
                    end_column=len(visible_columns)
                )

            for col_index in range(1, len(visible_columns) + 1):
                cell = worksheet.cell(row=excel_row_index, column=col_index)
                cell.border = border
                cell.alignment = Alignment(horizontal="right", vertical="center")

                if col_index in [1, 2]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                if row_type == "section_title":
                    cell.fill = section_fill
                    cell.font = black_bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif row_type == "facility_title":
                    cell.fill = facility_fill
                    cell.font = black_bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif row_type == "diff":
                    cell.fill = diff_fill
                    cell.font = black_bold_font

                    if col_index in [*month_col_indexes, total_col_index]:
                        value = cell.value

                        if isinstance(value, (int, float)):
                            if value > 0:
                                cell.font = red_font
                            elif value < 0:
                                cell.font = blue_font
                            else:
                                cell.font = black_bold_font

                elif row_type == "blank":
                    cell.fill = white_fill

                if col_index in [*month_col_indexes, total_col_index]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0.000'

            if row_type == "blank":
                worksheet.row_dimensions[excel_row_index].height = 8

        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["B"].width = 20

        for col_index in range(3, 15):
            col_letter = get_column_letter(col_index)
            worksheet.column_dimensions[col_letter].width = 12

        worksheet.column_dimensions["O"].width = 14
        worksheet.freeze_panes = "C2"

    output.seek(0)

    return output


def render_site_db_page():
    st.subheader("사업장별 DB")

    st.write(
        "선택한 사업장의 최근 3개년 월별 에너지 사용량과 온실가스 배출량을 시설별로 조회합니다. "
        "기준연도와 전년도 배출량의 증감량을 함께 표시합니다."
    )

    col1, col2 = st.columns(2)

    with col1:
        selected_site = st.selectbox(
            "사업장",
            list(SITES.keys())
        )

    with col2:
        selected_year = st.number_input(
            "기준연도",
            min_value=START_YEAR + 2,
            max_value=MAX_YEAR,
            value=max(get_default_year(), START_YEAR + 2),
            step=1
        )

    st.caption(
        f"조회 기준: {int(selected_year) - 2}년, {int(selected_year) - 1}년, {int(selected_year)}년"
    )

    site_db_df = build_site_db_table(
        site=selected_site,
        base_year=int(selected_year)
    )

    if site_db_df.empty:
        st.info("표시할 데이터가 없습니다.")
        return

    html_table = build_site_db_html_table(site_db_df)

    st.markdown(
        html_table,
        unsafe_allow_html=True
    )

    excel_file = make_site_db_excel_file(
        df=site_db_df,
        site=selected_site,
        base_year=int(selected_year)
    )

    st.download_button(
        label="사업장별 DB 엑셀 다운로드",
        data=excel_file,
        file_name=f"GHG_사업장별_DB_{selected_site}_{int(selected_year)}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def get_energy_monthly_total_emission_values(df, energy_type, year):
    if df.empty:
        return [0.0] * 12

    target_df = df[
        (df["energy_type"] == energy_type)
        & (df["year"] == year)
    ].copy()

    month_values = []

    for month in range(1, 13):
        month_sum = target_df[
            target_df["month"] == month
        ]["emission_amount"].sum()

        month_values.append(round(float(month_sum), 3))

    return month_values


def get_energy_site_monthly_values(df, energy_type, site, year):
    if df.empty:
        return {
            "usage_values": [0.0] * 12,
            "emission_values": [0.0] * 12,
            "usage_unit": "단위없음",
            "is_mixed_unit": False,
        }

    target_df = df[
        (df["energy_type"] == energy_type)
        & (df["site"] == site)
        & (df["year"] == year)
    ].copy()

    if target_df.empty:
        return {
            "usage_values": [0.0] * 12,
            "emission_values": [0.0] * 12,
            "usage_unit": "단위없음",
            "is_mixed_unit": False,
        }

    usage_units = (
        target_df["usage_unit"]
        .dropna()
        .astype(str)
        .str.strip()
        .replace("", pd.NA)
        .dropna()
        .unique()
        .tolist()
    )

    if len(usage_units) == 0:
        usage_unit = "단위없음"
        is_mixed_unit = False
    elif len(usage_units) == 1:
        usage_unit = usage_units[0]
        is_mixed_unit = False
    else:
        usage_unit = "단위혼합"
        is_mixed_unit = True

    usage_values = []
    emission_values = []

    for month in range(1, 13):
        month_df = target_df[target_df["month"] == month].copy()

        if is_mixed_unit:
            usage_sum = 0.0
        else:
            usage_sum = month_df["usage_amount"].sum()

        emission_sum = month_df["emission_amount"].sum()

        usage_values.append(round(float(usage_sum), 3))
        emission_values.append(round(float(emission_sum), 3))

    return {
        "usage_values": usage_values,
        "emission_values": emission_values,
        "usage_unit": usage_unit,
        "is_mixed_unit": is_mixed_unit,
    }


def calculate_change_rate_values_for_energy(previous_values, current_values):
    rate_values = []

    for previous_value, current_value in zip(previous_values, current_values):
        if previous_value == 0:
            rate_values.append("")
        else:
            change_rate = (current_value - previous_value) / previous_value * 100
            rate_values.append(round(float(change_rate), 1))

    return rate_values


def build_energy_db_table(energy_type, base_year):
    df = load_activity_data()

    year_1 = base_year - 2
    year_2 = base_year - 1
    year_3 = base_year

    years = [year_1, year_2, year_3]

    rows = []

    title_row = {
        "row_type": "section_title",
        "사업장": "온실가스 배출량",
        "구분": "",
    }

    for month_name in MONTH_COLUMNS:
        title_row[month_name] = ""

    title_row["합계"] = ""
    rows.append(title_row)

    yearly_total_emission_map = {}

    for year in years:
        emission_values = get_energy_monthly_total_emission_values(
            df=df,
            energy_type=energy_type,
            year=year,
        )

        yearly_total_emission_map[year] = emission_values

        row = {
            "row_type": "total_emission",
            "사업장": f"{str(year)[-2:]}년 실적",
            "구분": "배출량(tCO₂e)",
        }

        for index, month_name in enumerate(MONTH_COLUMNS):
            row[month_name] = emission_values[index]

        row["합계"] = round(sum(emission_values), 3)
        rows.append(row)

    previous_values = yearly_total_emission_map[year_2]
    current_values = yearly_total_emission_map[year_3]

    rate_values = calculate_change_rate_values_for_energy(
        previous_values=previous_values,
        current_values=current_values,
    )

    previous_total = sum(previous_values)
    current_total = sum(current_values)

    if previous_total == 0:
        total_rate = ""
    else:
        total_rate = round((current_total - previous_total) / previous_total * 100, 1)

    rate_row = {
        "row_type": "rate",
        "사업장": "실적 비교",
        "구분": "전년대비 증감(%)",
        "합계": total_rate,
    }

    for index, month_name in enumerate(MONTH_COLUMNS):
        rate_row[month_name] = rate_values[index]

    rows.append(rate_row)

    blank_row = {
        "row_type": "blank",
        "사업장": "",
        "구분": "",
    }

    for month_name in MONTH_COLUMNS:
        blank_row[month_name] = ""

    blank_row["합계"] = ""
    rows.append(blank_row)

    energy_df = df[df["energy_type"] == energy_type].copy()

    if not energy_df.empty:
        site_items = (
            energy_df[["site"]]
            .drop_duplicates()
            .sort_values(["site"])
        )

        site_order = {site: index for index, site in enumerate(SITES.keys())}
        site_items["site_order"] = site_items["site"].map(site_order)
        site_items = site_items.sort_values(["site_order", "site"]).drop(columns=["site_order"])

        for _, item in site_items.iterrows():
            site = str(item["site"])

            site_title = {
                "row_type": "site_title",
                "사업장": site,
                "구분": "",
            }

            for month_name in MONTH_COLUMNS:
                site_title[month_name] = ""

            site_title["합계"] = ""
            rows.append(site_title)

            year_emission_map = {}
            year_values_map = {}

            site_all_year_df = energy_df[energy_df["site"] == site].copy()

            site_units = (
                site_all_year_df["usage_unit"]
                .dropna()
                .astype(str)
                .str.strip()
                .replace("", pd.NA)
                .dropna()
                .unique()
                .tolist()
            )

            if len(site_units) == 0:
                display_usage_unit = "단위없음"
                display_is_mixed_unit = False
            elif len(site_units) == 1:
                display_usage_unit = site_units[0]
                display_is_mixed_unit = False
            else:
                display_usage_unit = "단위혼합"
                display_is_mixed_unit = True

            for year in years:
                values = get_energy_site_monthly_values(
                    df=df,
                    energy_type=energy_type,
                    site=site,
                    year=year,
                )

                usage_values = values["usage_values"]
                emission_values = values["emission_values"]

                if display_is_mixed_unit:
                    usage_values_for_display = ["-"] * 12
                    usage_total_for_display = "-"
                else:
                    usage_values_for_display = usage_values
                    usage_total_for_display = round(sum(usage_values), 3)

                year_emission_map[year] = emission_values
                year_values_map[year] = values

                usage_row = {
                    "row_type": "usage_mixed" if display_is_mixed_unit else "usage",
                    "사업장": f"{str(year)[-2:]}년 실적",
                    "구분": f"사용량({display_usage_unit})",
                }

                for index, month_name in enumerate(MONTH_COLUMNS):
                    usage_row[month_name] = usage_values_for_display[index]

                usage_row["합계"] = usage_total_for_display
                rows.append(usage_row)

                emission_row = {
                    "row_type": "emission",
                    "사업장": "",
                    "구분": "배출량(tCO₂e)",
                }

                for index, month_name in enumerate(MONTH_COLUMNS):
                    emission_row[month_name] = emission_values[index]

                emission_row["합계"] = round(sum(emission_values), 3)
                rows.append(emission_row)

            previous_site_values = year_emission_map[year_2]
            current_site_values = year_emission_map[year_3]

            site_diff_values = [
                round(current - previous, 3)
                for previous, current in zip(previous_site_values, current_site_values)
            ]

            site_diff_row = {
                "row_type": "diff",
                "사업장": "실적 비교",
                "구분": "전년대비 증감(tCO₂e)",
            }

            for index, month_name in enumerate(MONTH_COLUMNS):
                site_diff_row[month_name] = site_diff_values[index]

            site_diff_row["합계"] = round(
                sum(current_site_values) - sum(previous_site_values),
                3
            )

            rows.append(site_diff_row)

            site_blank_row = {
                "row_type": "blank",
                "사업장": "",
                "구분": "",
            }

            for month_name in MONTH_COLUMNS:
                site_blank_row[month_name] = ""

            site_blank_row["합계"] = ""
            rows.append(site_blank_row)

    result_df = pd.DataFrame(rows)

    display_columns = [
        "사업장",
        "구분",
        *MONTH_COLUMNS,
        "합계",
        "row_type",
    ]

    result_df = result_df[display_columns]

    return result_df


def format_energy_db_cell(value, is_rate_row=False):
    if value == "":
        return "-"

    if pd.isna(value):
        return "-"

    try:
        numeric_value = float(value)
    except Exception:
        return str(value)

    if numeric_value == 0:
        if is_rate_row:
            return "0.0"
        return "-"

    if is_rate_row:
        return f"{numeric_value:,.1f}"

    return f"{numeric_value:,.3f}"


def get_energy_value_class(value, is_rate_row=False):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return ""

    if numeric_value > 0:
        return " value-increase"

    if numeric_value < 0:
        return " value-decrease"

    return " value-zero"


def build_energy_db_html_table(df):
    table_columns = [
        "사업장",
        "구분",
        *MONTH_COLUMNS,
        "합계",
    ]

    html = """
    <style>
        .energy-db-table-wrapper {
            width: 100%;
            overflow-x: auto;
            margin-top: 12px;
            margin-bottom: 24px;
        }

        .energy-db-table {
            width: 100%;
            border-collapse: collapse;
            font-size: 12px;
            table-layout: fixed;
        }

        .energy-db-table th {
            background-color: #F3F4F6;
            color: #111827;
            border: 1px solid #D1D5DB;
            padding: 6px 5px;
            text-align: center;
            font-weight: 700;
            white-space: nowrap;
        }

        .energy-db-table td {
            border: 1px solid #D1D5DB;
            padding: 5px 5px;
            text-align: right;
            white-space: nowrap;
        }

        .energy-db-table th.col-site,
        .energy-db-table td.col-site {
            width: 170px;
            text-align: left;
        }

        .energy-db-table th.col-type,
        .energy-db-table td.col-type {
            width: 140px;
            text-align: center;
        }

        .energy-db-table th.col-month {
            width: 76px;
        }

        .energy-db-table th.col-total {
            width: 95px;
        }

        .energy-db-table tr.section-title-row td {
            background-color: #D9EAF7;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .energy-db-table tr.site-title-row td {
            background-color: #E5E7EB;
            color: #111827;
            font-weight: 800;
            text-align: left;
        }

        .energy-db-table tr.diff-row td,
        .energy-db-table tr.rate-row td {
            background-color: #FFF2CC;
            font-weight: 700;
        }

        .energy-db-table tr.blank-row td {
            background-color: #FFFFFF;
            height: 10px;
            border-left-color: #FFFFFF;
            border-right-color: #FFFFFF;
        }

        .energy-db-table td.value-increase {
            color: #B91C1C;
            font-weight: 800;
        }

        .energy-db-table td.value-decrease {
            color: #1D4ED8;
            font-weight: 800;
        }

        .energy-db-table td.value-zero {
            color: #111827;
            font-weight: 700;
        }
    </style>
    """

    html += '<div class="energy-db-table-wrapper">'
    html += '<table class="energy-db-table">'
    html += "<thead><tr>"

    for col in table_columns:
        class_name = ""

        if col == "사업장":
            class_name = "col-site"
        elif col == "구분":
            class_name = "col-type"
        elif col in MONTH_COLUMNS:
            class_name = "col-month"
        elif col == "합계":
            class_name = "col-total"

        html += f'<th class="{class_name}">{col}</th>'

    html += "</tr></thead>"
    html += "<tbody>"

    for _, row in df.iterrows():
        row_type = str(row["row_type"])

        row_classes = []

        if row_type == "section_title":
            row_classes.append("section-title-row")

        if row_type == "site_title":
            row_classes.append("site-title-row")

        if row_type == "diff":
            row_classes.append("diff-row")

        if row_type == "rate":
            row_classes.append("rate-row")

        if row_type == "blank":
            row_classes.append("blank-row")

        row_class_text = " ".join(row_classes)

        html += f'<tr class="{row_class_text}">'

        if row_type in ["section_title", "site_title"]:
            title_text = str(row["사업장"])

            html += f'<td class="col-site" colspan="{len(table_columns)}">{title_text}</td>'
            html += "</tr>"
            continue

        for col in table_columns:
            value = row[col]
            is_rate_or_diff_row = row_type in ["rate", "diff"]
            value_class = get_energy_value_class(value) if is_rate_or_diff_row and col not in ["사업장", "구분"] else ""

            if col == "사업장":
                html += f'<td class="col-site">{value}</td>'
            elif col == "구분":
                html += f'<td class="col-type">{value}</td>'
            else:
                html += f'<td class="{value_class.strip()}">{format_energy_db_cell(value, is_rate_row=(row_type == "rate"))}</td>'

        html += "</tr>"

    html += "</tbody></table></div>"

    return html


def make_energy_db_excel_file(df, energy_type, base_year):
    output = BytesIO()

    export_df = df.drop(columns=["row_type"]).copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(
            writer,
            index=False,
            sheet_name=f"{energy_type}_에너지원DB"
        )

        worksheet = writer.sheets[f"{energy_type}_에너지원DB"]

        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="F3F4F6")
        section_fill = PatternFill("solid", fgColor="D9EAF7")
        site_fill = PatternFill("solid", fgColor="E5E7EB")
        diff_fill = PatternFill("solid", fgColor="FFF2CC")
        white_fill = PatternFill("solid", fgColor="FFFFFF")

        border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        red_font = Font(color="B91C1C", bold=True)
        blue_font = Font(color="1D4ED8", bold=True)
        black_bold_font = Font(color="111827", bold=True)
        header_font = Font(color="111827", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        visible_columns = export_df.columns.tolist()
        month_col_indexes = [
            visible_columns.index(month_name) + 1
            for month_name in MONTH_COLUMNS
        ]
        total_col_index = visible_columns.index("합계") + 1

        for excel_row_index, row in enumerate(df.itertuples(index=False), start=2):
            row_dict = row._asdict()
            row_type = row_dict["row_type"]

            if row_type in ["section_title", "site_title"]:
                worksheet.merge_cells(
                    start_row=excel_row_index,
                    start_column=1,
                    end_row=excel_row_index,
                    end_column=len(visible_columns)
                )

            for col_index in range(1, len(visible_columns) + 1):
                cell = worksheet.cell(row=excel_row_index, column=col_index)
                cell.border = border
                cell.alignment = Alignment(horizontal="right", vertical="center")

                if col_index in [1, 2]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                if row_type == "section_title":
                    cell.fill = section_fill
                    cell.font = black_bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif row_type == "site_title":
                    cell.fill = site_fill
                    cell.font = black_bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                elif row_type in ["diff", "rate"]:
                    cell.fill = diff_fill
                    cell.font = black_bold_font

                    if col_index in [*month_col_indexes, total_col_index]:
                        value = cell.value

                        if isinstance(value, (int, float)):
                            if value > 0:
                                cell.font = red_font
                            elif value < 0:
                                cell.font = blue_font
                            else:
                                cell.font = black_bold_font

                elif row_type == "blank":
                    cell.fill = white_fill

                if col_index in [*month_col_indexes, total_col_index]:
                    if isinstance(cell.value, (int, float)):
                        if row_type == "rate":
                            cell.number_format = '#,##0.0'
                        else:
                            cell.number_format = '#,##0.000'

            if row_type == "blank":
                worksheet.row_dimensions[excel_row_index].height = 8

        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["B"].width = 20

        for col_index in range(3, 15):
            col_letter = get_column_letter(col_index)
            worksheet.column_dimensions[col_letter].width = 12

        worksheet.column_dimensions["O"].width = 14
        worksheet.freeze_panes = "C2"

    output.seek(0)

    return output


def render_energy_db_page():
    st.subheader("에너지원별 DB")

    st.write(
        "선택한 에너지원의 최근 3개년 월별 사용량과 온실가스 배출량을 사업장별로 조회합니다. "
        "총 배출량은 전년대비 증감률로, 사업장별 배출량은 전년대비 증감량으로 표시합니다."
    )

    col1, col2 = st.columns(2)

    with col1:
        selected_energy_type = st.selectbox(
            "에너지원",
            ENERGY_TYPES
        )

    with col2:
        selected_year = st.number_input(
            "기준연도",
            min_value=START_YEAR + 2,
            max_value=MAX_YEAR,
            value=max(get_default_year(), START_YEAR + 2),
            step=1
        )

    st.caption(
        f"조회 기준: {int(selected_year) - 2}년, {int(selected_year) - 1}년, {int(selected_year)}년"
    )

    energy_db_df = build_energy_db_table(
        energy_type=selected_energy_type,
        base_year=int(selected_year)
    )

    if energy_db_df.empty:
        st.info("표시할 데이터가 없습니다.")
        return

    html_table = build_energy_db_html_table(energy_db_df)

    st.markdown(
        html_table,
        unsafe_allow_html=True
    )

    excel_file = make_energy_db_excel_file(
        df=energy_db_df,
        energy_type=selected_energy_type,
        base_year=int(selected_year)
    )

    st.download_button(
        label="에너지원별 DB 엑셀 다운로드",
        data=excel_file,
        file_name=f"GHG_에너지원별_DB_{selected_energy_type}_{int(selected_year)}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def sync_activity_corporations():
    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    for site, site_info in SITES.items():
        cur.execute(
            """
            UPDATE activity_data
            SET corporation = ?,
                updated_at = ?
            WHERE site = ?
            """,
            (
                site_info["corporation"],
                now_text,
                site,
            ),
        )

    conn.commit()
    conn.close()

def load_reduction_targets():
    conn = get_connection()
    df = pd.read_sql_query(
        """
        SELECT
            id,
            year,
            corporation,
            base_emission,
            reduction_rate,
            target_emission,
            memo,
            created_at,
            updated_at
        FROM reduction_targets
        ORDER BY year, corporation
        """,
        conn,
    )
    conn.close()
    return df


def upsert_reduction_target(
    year,
    corporation,
    base_emission,
    reduction_rate,
    memo
):
    if corporation not in CORPORATIONS:
        raise ValueError("등록되지 않은 법인입니다.")

    if base_emission is None:
        base_emission = 0

    if reduction_rate is None:
        reduction_rate = 0

    base_emission = round(float(base_emission), 3)
    reduction_rate = round(float(reduction_rate), 3)

    if base_emission < 0:
        raise ValueError("기준배출량은 0 이상이어야 합니다.")

    target_emission = round(
        base_emission * (1 - reduction_rate / 100),
        3
    )

    if target_emission < 0:
        target_emission = 0

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        INSERT INTO reduction_targets
        (
            year,
            corporation,
            base_emission,
            reduction_rate,
            target_emission,
            memo,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(year, corporation)
        DO UPDATE SET
            base_emission = excluded.base_emission,
            reduction_rate = excluded.reduction_rate,
            target_emission = excluded.target_emission,
            memo = excluded.memo,
            updated_at = excluded.updated_at
        """,
        (
            year,
            corporation,
            base_emission,
            reduction_rate,
            target_emission,
            memo,
            now_text,
            now_text,
        ),
    )

    conn.commit()
    conn.close()


def build_reduction_target_template(year):
    target_df = load_reduction_targets()
    year_target_df = target_df[target_df["year"] == year].copy()

    activity_df = load_activity_data()

    previous_year = year - 1

    if activity_df.empty:
        previous_summary = pd.DataFrame(
            columns=[
                "corporation",
                "previous_emission",
            ]
        )
    else:
        previous_summary = (
            activity_df[
                (activity_df["year"] == previous_year)
                & (activity_df["corporation"].isin(CORPORATIONS))
            ]
            .groupby("corporation", as_index=False)["emission_amount"]
            .sum()
            .rename(columns={"emission_amount": "previous_emission"})
        )

    rows = []

    for corporation in CORPORATIONS:
        saved_target = year_target_df[
            year_target_df["corporation"] == corporation
        ].copy()

        previous_value_df = previous_summary[
            previous_summary["corporation"] == corporation
        ].copy()

        if not saved_target.empty:
            base_emission = round(float(saved_target.iloc[0]["base_emission"]), 3)
            reduction_rate = round(float(saved_target.iloc[0]["reduction_rate"]), 3)
            target_emission = round(float(saved_target.iloc[0]["target_emission"]), 3)
            memo = str(saved_target.iloc[0]["memo"])
        else:
            if not previous_value_df.empty:
                base_emission = round(float(previous_value_df.iloc[0]["previous_emission"]), 3)
            else:
                base_emission = 0.0

            reduction_rate = 0.0
            target_emission = round(base_emission, 3)
            memo = ""

        rows.append(
            {
                "corporation": corporation,
                "base_emission": base_emission,
                "reduction_rate": reduction_rate,
                "target_emission": target_emission,
                "memo": memo,
            }
        )

    return pd.DataFrame(rows)


def save_reduction_target_data(year, edited_df):
    saved_count = 0

    for _, row in edited_df.iterrows():
        corporation = str(row["corporation"]).strip()

        if corporation not in CORPORATIONS:
            continue

        base_emission = round(float(row["base_emission"]), 3)
        reduction_rate = round(float(row["reduction_rate"]), 3)
        memo = str(row["memo"]).strip()

        upsert_reduction_target(
            year=year,
            corporation=corporation,
            base_emission=base_emission,
            reduction_rate=reduction_rate,
            memo=memo,
        )

        saved_count += 1

    return saved_count

def sync_reduction_base_emission_from_previous_year(year):
    previous_year = year - 1

    activity_df = load_activity_data()
    target_df = load_reduction_targets()

    if activity_df.empty:
        return {
            "previous_year": previous_year,
            "synced_count": 0,
            "missing_count": len(CORPORATIONS),
            "message": "활동자료 데이터가 없습니다.",
        }

    previous_summary = (
        activity_df[
            (activity_df["year"] == previous_year)
            & (activity_df["corporation"].isin(CORPORATIONS))
        ]
        .groupby("corporation", as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "previous_emission"})
    )

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    synced_count = 0
    missing_count = 0

    for corporation in CORPORATIONS:
        previous_value_df = previous_summary[
            previous_summary["corporation"] == corporation
        ].copy()

        if previous_value_df.empty:
            base_emission = 0.0
            missing_count += 1
        else:
            base_emission = round(float(previous_value_df.iloc[0]["previous_emission"]), 3)

        saved_target = target_df[
            (target_df["year"] == year)
            & (target_df["corporation"] == corporation)
        ].copy()

        if saved_target.empty:
            reduction_rate = 0.0
            memo = ""
        else:
            reduction_rate = round(float(saved_target.iloc[0]["reduction_rate"]), 3)
            memo = str(saved_target.iloc[0]["memo"])

        target_emission = round(
            base_emission * (1 - reduction_rate / 100),
            3
        )

        if target_emission < 0:
            target_emission = 0

        cur.execute(
            """
            INSERT INTO reduction_targets
            (
                year,
                corporation,
                base_emission,
                reduction_rate,
                target_emission,
                memo,
                created_at,
                updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(year, corporation)
            DO UPDATE SET
                base_emission = excluded.base_emission,
                reduction_rate = excluded.reduction_rate,
                target_emission = excluded.target_emission,
                memo = excluded.memo,
                updated_at = excluded.updated_at
            """,
            (
                year,
                corporation,
                base_emission,
                reduction_rate,
                target_emission,
                memo,
                now_text,
                now_text,
            ),
        )

        synced_count += 1

    conn.commit()
    conn.close()

    return {
        "previous_year": previous_year,
        "synced_count": synced_count,
        "missing_count": missing_count,
        "message": "",
    }


def build_target_performance_table(year):
    target_df = load_reduction_targets()
    activity_df = load_activity_data()

    target_year_df = target_df[
        (target_df["year"] == year)
        & (target_df["corporation"].isin(CORPORATIONS))
    ].copy()

    if activity_df.empty:
        actual_summary = pd.DataFrame(
            columns=[
                "corporation",
                "actual_emission",
            ]
        )
    else:
        actual_summary = (
            activity_df[
                (activity_df["year"] == year)
                & (activity_df["corporation"].isin(CORPORATIONS))
            ]
            .groupby("corporation", as_index=False)["emission_amount"]
            .sum()
            .rename(columns={"emission_amount": "actual_emission"})
        )

    result_df = pd.DataFrame(
        {
            "corporation": CORPORATIONS
        }
    )

    result_df = pd.merge(
        result_df,
        target_year_df[
            [
                "corporation",
                "base_emission",
                "reduction_rate",
                "target_emission",
            ]
        ],
        on="corporation",
        how="left",
    )

    result_df = pd.merge(
        result_df,
        actual_summary,
        on="corporation",
        how="left",
    )

    result_df["base_emission"] = result_df["base_emission"].fillna(0)
    result_df["reduction_rate"] = result_df["reduction_rate"].fillna(0)
    result_df["target_emission"] = result_df["target_emission"].fillna(0)
    result_df["actual_emission"] = result_df["actual_emission"].fillna(0)

    result_df["gap_amount"] = (
        result_df["actual_emission"]
        - result_df["target_emission"]
    ).round(3)

    result_df["achievement_rate"] = result_df.apply(
        lambda row: round(row["actual_emission"] / row["target_emission"] * 100, 1)
        if row["target_emission"] > 0 else "",
        axis=1,
    )

    result_df = result_df.rename(
        columns={
            "corporation": "법인",
            "base_emission": "기준배출량(tCO₂e)",
            "reduction_rate": "감축목표율(%)",
            "target_emission": "목표배출량(tCO₂e)",
            "actual_emission": "실제배출량(tCO₂e)",
            "gap_amount": "목표 대비 차이(tCO₂e)",
            "achievement_rate": "목표 대비 배출률(%)",
        }
    )

    return result_df

def build_target_ytd_dashboard_table(selected_year, selected_month):
    target_df = load_reduction_targets()
    activity_df = load_activity_data()

    target_year_df = target_df[
        (target_df["year"] == selected_year)
        & (target_df["corporation"].isin(CORPORATIONS))
    ].copy()

    if activity_df.empty:
        actual_ytd_summary = pd.DataFrame(
            columns=[
                "corporation",
                "actual_ytd_emission",
            ]
        )
    else:
        actual_ytd_summary = (
            activity_df[
                (activity_df["year"] == selected_year)
                & (activity_df["month"] <= selected_month)
                & (activity_df["corporation"].isin(CORPORATIONS))
            ]
            .groupby("corporation", as_index=False)["emission_amount"]
            .sum()
            .rename(columns={"emission_amount": "actual_ytd_emission"})
        )

    result_df = pd.DataFrame(
        {
            "corporation": CORPORATIONS
        }
    )

    if target_year_df.empty:
        result_df["annual_target_emission"] = 0.0
    else:
        result_df = pd.merge(
            result_df,
            target_year_df[
                [
                    "corporation",
                    "target_emission",
                ]
            ],
            on="corporation",
            how="left",
        )

        result_df = result_df.rename(
            columns={
                "target_emission": "annual_target_emission"
            }
        )

    result_df = pd.merge(
        result_df,
        actual_ytd_summary,
        on="corporation",
        how="left",
    )

    result_df["annual_target_emission"] = result_df["annual_target_emission"].fillna(0)
    result_df["actual_ytd_emission"] = result_df["actual_ytd_emission"].fillna(0)

    result_df["ytd_target_emission"] = (
        result_df["annual_target_emission"] * selected_month / 12
    ).round(3)

    result_df["target_gap"] = (
        result_df["actual_ytd_emission"] - result_df["ytd_target_emission"]
    ).round(3)

    result_df["target_usage_rate"] = result_df.apply(
        lambda row: round(row["actual_ytd_emission"] / row["ytd_target_emission"] * 100, 1)
        if row["ytd_target_emission"] > 0 else "",
        axis=1,
    )

    result_df = result_df.rename(
        columns={
            "corporation": "법인",
            "annual_target_emission": "연간 목표배출량(tCO₂e)",
            "ytd_target_emission": "누적 목표배출량(tCO₂e)",
            "actual_ytd_emission": "누적 실제배출량(tCO₂e)",
            "target_gap": "누적 목표 대비 차이(tCO₂e)",
            "target_usage_rate": "누적 목표 대비 배출률(%)",
        }
    )

    return result_df

def load_revenue_data():
    conn = get_connection()
    df = pd.read_sql_query(
        """
        SELECT
            id,
            year,
            month,
            corporation,
            revenue_okr,
            created_at,
            updated_at
        FROM monthly_revenue
        ORDER BY year, month, corporation
        """,
        conn,
    )
    conn.close()
    return df


def upsert_revenue_data(year, month, corporation, revenue_okr):
    if corporation not in CORPORATIONS:
        raise ValueError("등록되지 않은 법인입니다.")

    if revenue_okr is None:
        revenue_okr = 0

    revenue_okr = round(float(revenue_okr), 3)

    if revenue_okr < 0:
        raise ValueError("매출액은 0 이상이어야 합니다.")

    now_text = datetime.now().isoformat(timespec="seconds")

    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        INSERT INTO monthly_revenue
        (
            year,
            month,
            corporation,
            revenue_okr,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(year, month, corporation)
        DO UPDATE SET
            revenue_okr = excluded.revenue_okr,
            updated_at = excluded.updated_at
        """,
        (
            year,
            month,
            corporation,
            revenue_okr,
            now_text,
            now_text,
        ),
    )

    conn.commit()
    conn.close()


def build_revenue_input_template(year):
    revenue_df = load_revenue_data()
    year_df = revenue_df[revenue_df["year"] == year].copy()

    rows = []

    for corporation in CORPORATIONS:
        row_data = {
            "corporation": corporation,
        }

        target_df = year_df[
            year_df["corporation"] == corporation
        ].copy()

        for month in range(1, 13):
            month_name = f"{month}월"

            month_values = target_df[
                target_df["month"] == month
            ]["revenue_okr"]

            if not month_values.empty:
                row_data[month_name] = round(float(month_values.iloc[0]), 3)
            else:
                row_data[month_name] = 0.0

        rows.append(row_data)

    result_df = pd.DataFrame(rows)

    return result_df


def save_revenue_input_data(year, edited_df):
    saved_count = 0

    for _, row in edited_df.iterrows():
        corporation = str(row["corporation"]).strip()

        if corporation not in CORPORATIONS:
            continue

        for month in range(1, 13):
            month_name = f"{month}월"
            revenue_okr = round(float(row[month_name]), 3)

            upsert_revenue_data(
                year=year,
                month=month,
                corporation=corporation,
                revenue_okr=revenue_okr,
            )

            saved_count += 1

    return saved_count


def make_revenue_excel_template(year):
    template_df = build_revenue_input_template(year)

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        template_df.to_excel(
            writer,
            index=False,
            sheet_name="매출액"
        )

        guide_df = pd.DataFrame(
            {
                "항목": [
                    "입력 연도",
                    "단위",
                    "작성 방법",
                    "주의 사항 1",
                    "주의 사항 2",
                ],
                "내용": [
                    year,
                    "억원",
                    "매출액 시트에서 법인별 1월~12월 매출액을 입력한 뒤 업로드하세요.",
                    "corporation 컬럼명과 법인명은 변경하지 마세요.",
                    "매출액은 0 이상의 숫자로 입력하세요. 소수점 셋째 자리까지 저장됩니다.",
                ],
            }
        )

        guide_df.to_excel(
            writer,
            index=False,
            sheet_name="작성가이드"
        )

        for sheet_name in ["매출액", "작성가이드"]:
            worksheet = writer.sheets[sheet_name]

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter

                for cell in column_cells:
                    cell_value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(cell_value))

                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output.seek(0)

    return output


def read_uploaded_revenue_excel(uploaded_file):
    try:
        uploaded_df = pd.read_excel(
            uploaded_file,
            sheet_name="매출액",
            engine="openpyxl"
        )
    except ValueError:
        uploaded_df = pd.read_excel(
            uploaded_file,
            engine="openpyxl"
        )

    uploaded_df.columns = [str(col).strip() for col in uploaded_df.columns]

    return uploaded_df


def validate_revenue_upload_df(uploaded_df):
    errors = []
    warnings = []

    required_columns = [
        "corporation",
        *MONTH_COLUMNS,
    ]

    missing_columns = [
        col for col in required_columns
        if col not in uploaded_df.columns
    ]

    if missing_columns:
        errors.append(
            "필수 컬럼이 누락되었습니다: "
            + ", ".join(missing_columns)
        )
        return errors, warnings

    if uploaded_df.empty:
        errors.append("업로드 파일에 데이터 행이 없습니다.")
        return errors, warnings

    duplicated_df = uploaded_df[
        uploaded_df.duplicated(
            subset=["corporation"],
            keep=False
        )
    ]

    if not duplicated_df.empty:
        duplicated_items = (
            duplicated_df["corporation"].astype(str)
        ).drop_duplicates().tolist()

        errors.append(
            "업로드 파일 안에 같은 법인이 중복되어 있습니다: "
            + ", ".join(duplicated_items)
        )

    for index, row in uploaded_df.iterrows():
        excel_row_number = index + 2

        corporation = str(row["corporation"]).strip()

        if corporation not in CORPORATIONS:
            errors.append(
                f"{excel_row_number}행: 등록되지 않은 법인입니다. corporation={corporation}"
            )

        for month_name in MONTH_COLUMNS:
            value = row[month_name]

            if pd.isna(value):
                warnings.append(
                    f"{excel_row_number}행 {month_name}: 빈 값은 0으로 처리됩니다."
                )
                continue

            try:
                numeric_value = float(value)
            except Exception:
                errors.append(
                    f"{excel_row_number}행 {month_name}: 숫자가 아닙니다. 입력값={value}"
                )
                continue

            if numeric_value < 0:
                errors.append(
                    f"{excel_row_number}행 {month_name}: 매출액은 0 이상이어야 합니다. 입력값={value}"
                )

    return errors, warnings


def normalize_revenue_upload_df(uploaded_df):
    result_df = uploaded_df.copy()

    result_df["corporation"] = result_df["corporation"].astype(str).str.strip()

    for month_name in MONTH_COLUMNS:
        result_df[month_name] = (
            pd.to_numeric(result_df[month_name], errors="coerce")
            .fillna(0)
            .round(3)
        )

    return result_df


def render_revenue_page():
    st.subheader("매출액 관리")

    st.write(
        "법인별 월별 매출액을 관리합니다. "
        "입력 단위는 억원이며, 대시보드의 탄소집약도 계산에 사용됩니다."
    )

    selected_year = st.number_input(
        "입력 연도",
        min_value=START_YEAR,
        max_value=MAX_YEAR,
        value=get_default_year(),
        step=1
    )

    template_df = build_revenue_input_template(
        year=int(selected_year)
    )

    st.subheader("법인별 월별 매출액 입력")

    month_column_config = {
        month_name: st.column_config.NumberColumn(
            month_name,
            min_value=0.0,
            step=0.001,
            format="%.3f"
        )
        for month_name in MONTH_COLUMNS
    }

    column_config = {
        "corporation": st.column_config.TextColumn("법인"),
    }
    column_config.update(month_column_config)

    edited_df = st.data_editor(
        template_df,
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        key="revenue_editor",
        disabled=[
            "corporation",
        ],
        column_config=column_config
    )

    if st.button("매출액 저장", type="primary"):
        try:
            saved_count = save_revenue_input_data(
                year=int(selected_year),
                edited_df=edited_df,
            )

            st.success(
                f"{int(selected_year)}년 월별 매출액이 저장되었습니다. "
                f"저장 또는 업데이트된 데이터 수: {saved_count}건"
            )
            st.rerun()

        except Exception as e:
            st.error(f"매출액 저장 중 오류가 발생했습니다: {e}")

    st.divider()

    st.subheader("저장된 매출액 데이터")

    revenue_df = load_revenue_data()

    if revenue_df.empty:
        st.info("저장된 매출액 데이터가 없습니다.")
    else:
        view_df = revenue_df[
            (revenue_df["year"] == int(selected_year))
            & (revenue_df["corporation"].isin(CORPORATIONS))
        ].copy()

        if view_df.empty:
            st.info(f"{int(selected_year)}년에 저장된 매출액 데이터가 없습니다.")
        else:
            display_df = view_df[
                [
                    "year",
                    "month",
                    "corporation",
                    "revenue_okr",
                    "updated_at",
                ]
            ].copy()

            display_df["revenue_okr"] = display_df["revenue_okr"].round(3)

            st.dataframe(
                display_df,
                use_container_width=True,
                hide_index=True,
            )

def render_reduction_target_page():
    st.subheader("감축목표 관리")

    st.write(
        "법인별 연간 감축목표를 관리합니다. "
        "기준배출량과 감축목표율을 입력하면 목표배출량이 자동 계산됩니다."
    )

    selected_year = st.number_input(
        "목표 연도",
        min_value=START_YEAR,
        max_value=MAX_YEAR,
        value=get_default_year(),
        step=1,
        key="target_year_input"
    )

    template_df = build_reduction_target_template(
        year=int(selected_year)
    )

    st.caption("기준배출량은 목표연도 전년도 실제 배출량입니다.")

    confirm_sync_base = st.checkbox(
        "기준배출량 갱신 확인",
        key="confirm_sync_reduction_base"
    )

    if st.button(
        "전년도 기준배출량 불러오기",
        disabled=not confirm_sync_base
    ):
        result = sync_reduction_base_emission_from_previous_year(
            year=int(selected_year)
        )

        st.success(
            f"{int(selected_year)}년 기준배출량 갱신이 완료되었습니다. "
            f"기준연도: {result['previous_year']}년, "
            f"처리 법인 수: {result['synced_count']}개, "
            f"전년도 데이터 없는 법인 수: {result['missing_count']}개"
        )

        st.rerun()

    st.subheader("법인별 감축목표 입력")

    edited_df = st.data_editor(
        template_df,
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        key="reduction_target_editor",
        disabled=[
            "corporation",
            "target_emission",
        ],
        column_config={
            "corporation": st.column_config.TextColumn("법인"),
            "base_emission": st.column_config.NumberColumn(
                "기준배출량(tCO₂e)",
                min_value=0.0,
                step=0.001,
                format="%.3f",
            ),
            "reduction_rate": st.column_config.NumberColumn(
                "감축목표율(%)",
                min_value=0.0,
                max_value=100.0,
                step=0.1,
                format="%.1f",
            ),
            "target_emission": st.column_config.NumberColumn(
                "목표배출량(tCO₂e)",
                min_value=0.0,
                step=0.001,
                format="%.3f",
            ),
            "memo": st.column_config.TextColumn("메모"),
        },
    )

    st.info(
        "화면의 목표배출량은 저장 후 다시 계산되어 반영됩니다. "
        "감축목표율을 수정한 뒤 저장하면 목표배출량이 업데이트됩니다."
    )

    if st.button("감축목표 저장", type="primary"):
        try:
            saved_count = save_reduction_target_data(
                year=int(selected_year),
                edited_df=edited_df,
            )

            st.success(
                f"{int(selected_year)}년 감축목표가 저장되었습니다. "
                f"저장 또는 업데이트된 데이터 수: {saved_count}건"
            )
            st.rerun()

        except Exception as e:
            st.error(f"감축목표 저장 중 오류가 발생했습니다: {e}")

    st.divider()

    st.subheader("목표 대비 실적")

    performance_df = build_target_performance_table(
        year=int(selected_year)
    )

    if performance_df.empty:
        st.info("목표 대비 실적 데이터가 없습니다.")
    else:
        st.dataframe(
            style_dashboard_table(performance_df),
            use_container_width=True,
            hide_index=True,
        )


def safe_change_rate(current_value, previous_value):
    if previous_value == 0:
        return ""

    return round((current_value - previous_value) / previous_value * 100, 1)


def format_dashboard_number(value, decimals=3, zero_as_dash=False):
    if value == "":
        return ""

    if pd.isna(value):
        return ""

    try:
        numeric_value = float(value)
    except Exception:
        return str(value)

    if zero_as_dash and numeric_value == 0:
        return "-"

    return f"{numeric_value:,.{decimals}f}"


def get_current_previous_month_df(df, selected_year, selected_month):
    current_df = df[
        (df["year"] == selected_year)
        & (df["month"] == selected_month)
    ].copy()

    previous_df = df[
        (df["year"] == selected_year - 1)
        & (df["month"] == selected_month)
    ].copy()

    return current_df, previous_df


def get_ytd_df(df, selected_year, selected_month):
    current_ytd_df = df[
        (df["year"] == selected_year)
        & (df["month"] <= selected_month)
    ].copy()

    previous_ytd_df = df[
        (df["year"] == selected_year - 1)
        & (df["month"] <= selected_month)
    ].copy()

    return current_ytd_df, previous_ytd_df


def get_revenue_for_month(selected_year, selected_month):
    revenue_df = load_revenue_data()

    if revenue_df.empty:
        return pd.DataFrame(
            columns=[
                "corporation",
                "revenue_okr",
            ]
        )

    target_df = revenue_df[
        (revenue_df["year"] == selected_year)
        & (revenue_df["month"] == selected_month)
        & (revenue_df["corporation"].isin(CORPORATIONS))
    ].copy()

    if target_df.empty:
        return pd.DataFrame(
            columns=[
                "corporation",
                "revenue_okr",
            ]
        )

    result_df = (
        target_df
        .groupby("corporation", as_index=False)["revenue_okr"]
        .sum()
    )

    return result_df


def build_corporation_dashboard_table(current_df, previous_df, selected_year, selected_month):
    current_summary = (
        current_df
        .groupby("corporation", as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "current_emission"})
    )

    previous_summary = (
        previous_df
        .groupby("corporation", as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "previous_emission"})
    )

    revenue_summary = get_revenue_for_month(
        selected_year=selected_year,
        selected_month=selected_month,
    )

    all_corporations = sorted(
        list(
            set(current_summary["corporation"].tolist())
            | set(previous_summary["corporation"].tolist())
            | set(CORPORATIONS)
        )
    )

    result_df = pd.DataFrame(
        {
            "corporation": all_corporations
        }
    )

    result_df = pd.merge(
        result_df,
        current_summary,
        on="corporation",
        how="left"
    )

    result_df = pd.merge(
        result_df,
        previous_summary,
        on="corporation",
        how="left"
    )

    result_df = pd.merge(
        result_df,
        revenue_summary,
        on="corporation",
        how="left"
    )

    result_df["current_emission"] = result_df["current_emission"].fillna(0)
    result_df["previous_emission"] = result_df["previous_emission"].fillna(0)
    result_df["revenue_okr"] = result_df["revenue_okr"].fillna(0)

    result_df["change_amount"] = (
        result_df["current_emission"]
        - result_df["previous_emission"]
    )

    result_df["change_rate"] = result_df.apply(
        lambda row: safe_change_rate(
            current_value=row["current_emission"],
            previous_value=row["previous_emission"],
        ),
        axis=1,
    )

    result_df["carbon_intensity"] = result_df.apply(
        lambda row: round(row["current_emission"] / row["revenue_okr"], 3)
        if row["revenue_okr"] > 0 and row["corporation"] in CORPORATIONS
        else "",
        axis=1,
    )

    corporation_order = {
        "예산법인": 1,
        "안성법인": 2,
        "멕시코법인": 3,
        "베트남법인": 4,
        "중국법인": 5,
        "비매출사업장": 6,
    }

    result_df["sort_order"] = result_df["corporation"].map(corporation_order).fillna(99)
    result_df = result_df.sort_values(["sort_order", "corporation"]).drop(columns=["sort_order"])

    result_df = result_df.rename(
        columns={
            "corporation": "법인",
            "current_emission": "당월 배출량(tCO₂e)",
            "previous_emission": "전년 동월 배출량(tCO₂e)",
            "change_amount": "증감량(tCO₂e)",
            "change_rate": "증감률(%)",
            "revenue_okr": "매출액(억원)",
            "carbon_intensity": "탄소집약도(tCO₂e/억원)",
        }
    )

    return result_df


def build_energy_dashboard_table(current_df, previous_df):
    current_summary = (
        current_df
        .groupby("energy_type", as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "current_emission"})
    )

    previous_summary = (
        previous_df
        .groupby("energy_type", as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "previous_emission"})
    )

    total_current = current_summary["current_emission"].sum()

    all_energy_types = sorted(
        list(
            set(current_summary["energy_type"].tolist())
            | set(previous_summary["energy_type"].tolist())
        )
    )

    result_df = pd.DataFrame(
        {
            "energy_type": all_energy_types
        }
    )

    result_df = pd.merge(
        result_df,
        current_summary,
        on="energy_type",
        how="left"
    )

    result_df = pd.merge(
        result_df,
        previous_summary,
        on="energy_type",
        how="left"
    )

    result_df["current_emission"] = result_df["current_emission"].fillna(0)
    result_df["previous_emission"] = result_df["previous_emission"].fillna(0)

    result_df["share"] = result_df["current_emission"].apply(
        lambda value: round(value / total_current * 100, 1)
        if total_current > 0 else 0
    )

    result_df["change_rate"] = result_df.apply(
        lambda row: safe_change_rate(
            current_value=row["current_emission"],
            previous_value=row["previous_emission"],
        ),
        axis=1,
    )

    energy_order = {
        energy_type: index
        for index, energy_type in enumerate(ENERGY_TYPES)
    }

    result_df["sort_order"] = result_df["energy_type"].map(energy_order).fillna(99)
    result_df = result_df.sort_values(["current_emission"], ascending=False)
    result_df = result_df.drop(columns=["sort_order"])

    result_df = result_df.rename(
        columns={
            "energy_type": "에너지원",
            "current_emission": "당월 배출량(tCO₂e)",
            "share": "비중(%)",
            "previous_emission": "전년 동월 배출량(tCO₂e)",
            "change_rate": "증감률(%)",
        }
    )

    return result_df


def build_top5_facility_table(current_df):
    if current_df.empty:
        return pd.DataFrame(
            columns=[
                "순위",
                "사업장",
                "시설",
                "에너지원",
                "배출량(tCO₂e)",
                "비중(%)",
            ]
        )

    total_emission = current_df["emission_amount"].sum()

    top_df = (
        current_df
        .groupby(["site", "facility", "energy_type"], as_index=False)["emission_amount"]
        .sum()
        .sort_values("emission_amount", ascending=False)
        .head(5)
    )

    top_df["비중(%)"] = top_df["emission_amount"].apply(
        lambda value: round(value / total_emission * 100, 1)
        if total_emission > 0 else 0
    )

    top_df.insert(0, "순위", range(1, len(top_df) + 1))

    top_df = top_df.rename(
        columns={
            "site": "사업장",
            "facility": "시설",
            "energy_type": "에너지원",
            "emission_amount": "배출량(tCO₂e)",
        }
    )

    return top_df


def build_top5_increase_facility_table(current_df, previous_df):
    current_summary = (
        current_df
        .groupby(["site", "facility", "energy_type"], as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "current_emission"})
    )

    previous_summary = (
        previous_df
        .groupby(["site", "facility", "energy_type"], as_index=False)["emission_amount"]
        .sum()
        .rename(columns={"emission_amount": "previous_emission"})
    )

    compare_df = pd.merge(
        current_summary,
        previous_summary,
        on=["site", "facility", "energy_type"],
        how="left"
    )

    if compare_df.empty:
        return pd.DataFrame(
            columns=[
                "순위",
                "사업장",
                "시설",
                "에너지원",
                "전년 동월(tCO₂e)",
                "당월(tCO₂e)",
                "증감량(tCO₂e)",
                "증감률(%)",
            ]
        )

    compare_df["previous_emission"] = compare_df["previous_emission"].fillna(0)
    compare_df["change_amount"] = (
        compare_df["current_emission"]
        - compare_df["previous_emission"]
    )

    compare_df["change_rate"] = compare_df.apply(
        lambda row: safe_change_rate(
            current_value=row["current_emission"],
            previous_value=row["previous_emission"],
        ),
        axis=1,
    )

    compare_df = (
        compare_df
        .sort_values("change_amount", ascending=False)
        .head(5)
    )

    compare_df.insert(0, "순위", range(1, len(compare_df) + 1))

    compare_df = compare_df.rename(
        columns={
            "site": "사업장",
            "facility": "시설",
            "energy_type": "에너지원",
            "previous_emission": "전년 동월(tCO₂e)",
            "current_emission": "당월(tCO₂e)",
            "change_amount": "증감량(tCO₂e)",
            "change_rate": "증감률(%)",
        }
    )

    return compare_df


def build_missing_facility_table(df, selected_year, selected_month):
    fe_df = load_facility_energy()

    if fe_df.empty:
        return pd.DataFrame(
            columns=[
                "사업장",
                "시설",
                "에너지원",
                "단위",
                "연소구분",
            ]
        )

    current_keys_df = df[
        (df["year"] == selected_year)
        & (df["month"] == selected_month)
    ][["site", "facility"]].drop_duplicates()

    merged_df = pd.merge(
        fe_df,
        current_keys_df,
        on=["site", "facility"],
        how="left",
        indicator=True,
    )

    missing_df = merged_df[
        merged_df["_merge"] == "left_only"
    ].copy()

    missing_df = missing_df[
        [
            "site",
            "facility",
            "energy_type",
            "usage_unit",
            "combustion_type",
        ]
    ].rename(
        columns={
            "site": "사업장",
            "facility": "시설",
            "energy_type": "에너지원",
            "usage_unit": "단위",
            "combustion_type": "연소구분",
        }
    )

    site_order = {site: index for index, site in enumerate(SITES.keys())}
    missing_df["site_order"] = missing_df["사업장"].map(site_order).fillna(99)
    missing_df = missing_df.sort_values(["site_order", "시설"]).drop(columns=["site_order"])

    return missing_df


def style_dashboard_table(df):
    numeric_columns = [
        col for col in df.columns
        if col != "순위"
        and col not in ["법인", "에너지원", "사업장", "시설", "단위", "연소구분"]
    ]

    format_dict = {}

    for col in numeric_columns:
        if "증감률" in col or "비중" in col:
            format_dict[col] = lambda value: "" if value == "" else f"{float(value):,.1f}" if isinstance(value, (int, float)) else value
        elif "매출액" in col:
            format_dict[col] = lambda value: "-" if value == "" or pd.isna(value) or float(value) == 0 else f"{float(value):,.3f}"
        elif "탄소집약도" in col:
            format_dict[col] = lambda value: "-" if value == "" or pd.isna(value) else f"{float(value):,.3f}"
        else:
            format_dict[col] = lambda value: "-" if value == "" or pd.isna(value) or float(value) == 0 else f"{float(value):,.3f}"

    def highlight_change(value):
        try:
            numeric_value = float(value)
        except Exception:
            return ""

        if numeric_value > 0:
            return "color: #B91C1C; font-weight: bold;"

        if numeric_value < 0:
            return "color: #1D4ED8; font-weight: bold;"

        return ""

    styler = df.style.format(format_dict)

    for col in df.columns:
        if "증감" in col:
            styler = styler.map(highlight_change, subset=[col])

    return styler


def make_dashboard(df, selected_year, selected_month):
    st.subheader("대시보드")

    if df.empty:
        st.info("입력된 데이터가 없습니다.")
        return

    current_df, previous_df = get_current_previous_month_df(
        df=df,
        selected_year=selected_year,
        selected_month=selected_month,
    )

    current_ytd_df, previous_ytd_df = get_ytd_df(
        df=df,
        selected_year=selected_year,
        selected_month=selected_month,
    )

    if current_df.empty:
        st.warning("선택한 월의 데이터가 없습니다.")
        return

    total_emission = current_df["emission_amount"].sum()
    previous_total = previous_df["emission_amount"].sum() if not previous_df.empty else 0

    yoy_change_amount = total_emission - previous_total
    yoy_change_rate = safe_change_rate(
        current_value=total_emission,
        previous_value=previous_total,
    )

    ytd_total = current_ytd_df["emission_amount"].sum()
    previous_ytd_total = previous_ytd_df["emission_amount"].sum() if not previous_ytd_df.empty else 0

    ytd_change_rate = safe_change_rate(
        current_value=ytd_total,
        previous_value=previous_ytd_total,
    )

    non_revenue_emission = current_df[
        current_df["corporation"] == "비매출사업장"
    ]["emission_amount"].sum()

    fe_df = load_facility_energy()
    registered_facility_count = len(fe_df)
    input_facility_count = len(
        current_df[["site", "facility"]].drop_duplicates()
    )

    if registered_facility_count > 0:
        input_rate = input_facility_count / registered_facility_count * 100
    else:
        input_rate = 0

    st.caption(
        f"조회 기준: {selected_year}년 {selected_month}월 "
        f"/ 전년 동월: {selected_year - 1}년 {selected_month}월"
    )

    metric_col1, metric_col2, metric_col3 = st.columns(3)

    metric_col1.metric(
        "당월 배출량",
        f"{total_emission:,.3f} tCO₂e",
        delta=f"{yoy_change_amount:,.3f} tCO₂e"
    )

    if yoy_change_rate == "":
        yoy_rate_text = "비교 불가"
    else:
        yoy_rate_text = f"{yoy_change_rate:,.1f}%"

    metric_col2.metric(
        "전년 동월 대비",
        yoy_rate_text
    )

    metric_col3.metric(
        "연간 누적 배출량",
        f"{ytd_total:,.3f} tCO₂e",
        delta=f"{ytd_total - previous_ytd_total:,.3f} tCO₂e"
    )

    metric_col4, metric_col5, metric_col6 = st.columns(3)

    if ytd_change_rate == "":
        ytd_rate_text = "비교 불가"
    else:
        ytd_rate_text = f"{ytd_change_rate:,.1f}%"

    metric_col4.metric(
        "전년 누적 대비",
        ytd_rate_text
    )

    metric_col5.metric(
        "데이터 입력률",
        f"{input_rate:,.1f}%",
        delta=f"{input_facility_count}/{registered_facility_count}개 시설"
    )

    metric_col6.metric(
        "비매출(과천 빌딩) 배출량",
        f"{non_revenue_emission:,.3f} tCO₂e"
    )

    st.divider()

    target_ytd_table = build_target_ytd_dashboard_table(
        selected_year=selected_year,
        selected_month=selected_month,
    )

    total_ytd_target = target_ytd_table["누적 목표배출량(tCO₂e)"].sum()
    total_ytd_actual = target_ytd_table["누적 실제배출량(tCO₂e)"].sum()

    total_ytd_gap = total_ytd_actual - total_ytd_target

    if total_ytd_target > 0:
        total_target_usage_rate = round(total_ytd_actual / total_ytd_target * 100, 1)
    else:
        total_target_usage_rate = ""

    exceeded_count = len(
        target_ytd_table[
            target_ytd_table["누적 목표 대비 차이(tCO₂e)"] > 0
        ]
    )

    st.subheader("감축목표 진행 현황")

    st.caption(
        "누적 목표배출량은 연간 목표배출량을 12개월로 균등 배분해 계산합니다."
    )

    target_col1, target_col2, target_col3, target_col4 = st.columns(4)

    with target_col1:
        st.metric(
            "누적 목표배출량",
            f"{total_ytd_target:,.3f} tCO₂e"
        )

    with target_col2:
        st.metric(
            "누적 실제배출량",
            f"{total_ytd_actual:,.3f} tCO₂e"
        )

    with target_col3:
        st.metric(
            "누적 목표 대비 차이",
            f"{total_ytd_gap:,.3f} tCO₂e"
        )

    with target_col4:
        if total_target_usage_rate == "":
            usage_rate_text = "목표 없음"
        else:
            usage_rate_text = f"{total_target_usage_rate:,.1f}%"

        st.metric(
            "누적 목표 대비 배출률",
            usage_rate_text,
            delta=f"목표 초과 법인 {exceeded_count}개"
        )

    st.dataframe(
        style_dashboard_table(target_ytd_table),
        use_container_width=True,
        hide_index=True,
    )

    st.divider()

    corporation_table = build_corporation_dashboard_table(
        current_df=current_df,
        previous_df=previous_df,
        selected_year=selected_year,
        selected_month=selected_month,
    )

    st.subheader("법인별 월간 실적")

    st.dataframe(
        style_dashboard_table(corporation_table),
        use_container_width=True,
        hide_index=True,
    )

    st.divider()

    energy_table = build_energy_dashboard_table(
        current_df=current_df,
        previous_df=previous_df,
    )

    st.subheader("에너지원별 월간 배출량")

    if energy_table.empty:
        st.info("에너지원별 배출량 데이터가 없습니다.")
    else:
        st.dataframe(
            style_dashboard_table(energy_table),
            use_container_width=True,
            hide_index=True,
        )

    st.divider()

    top_col1, top_col2 = st.columns(2)

    with top_col1:
        st.subheader("당월 배출량 TOP5")

        top5_table = build_top5_facility_table(
            current_df=current_df
        )

        st.dataframe(
            style_dashboard_table(top5_table),
            use_container_width=True,
            hide_index=True,
        )

    with top_col2:
        st.subheader("전년 동월 대비 증가 시설")

        increase_top5_table = build_top5_increase_facility_table(
            current_df=current_df,
            previous_df=previous_df,
        )

        st.dataframe(
            style_dashboard_table(increase_top5_table),
            use_container_width=True,
            hide_index=True,
        )

    st.divider()

    st.subheader("데이터 품질 점검")

    missing_df = build_missing_facility_table(
        df=df,
        selected_year=selected_year,
        selected_month=selected_month,
    )

    if missing_df.empty:
        st.success("선택한 월의 모든 등록 시설에 사용량 데이터가 입력되어 있습니다.")
    else:
        st.warning(f"선택한 월에 사용량이 입력되지 않은 시설이 {len(missing_df)}개 있습니다.")

        with st.expander("미입력 시설 목록 보기", expanded=False):
            st.dataframe(
                missing_df,
                use_container_width=True,
                hide_index=True,
            )


def render_dashboard_page():
    st.sidebar.subheader("조회 조건")

    st.write(
        "선택한 월의 배출량, 전년 대비 증감, 감축목표 진행 현황, "
        "법인별·에너지원별 현황을 요약합니다."
    )

    selected_year = st.sidebar.number_input(
        "연도",
        min_value=START_YEAR,
        max_value=MAX_YEAR,
        value=get_default_year(),
        step=1
    )

    selected_month = st.sidebar.selectbox(
        "월",
        list(range(1, 13)),
        index=datetime.now().month - 1
    )

    df = load_activity_data()
    make_dashboard(df, int(selected_year), int(selected_month))

def get_timestamp_text():
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def get_db_file_info():
    info = {
        "exists": os.path.exists(DB_NAME),
        "file_size_mb": 0,
        "modified_at": "",
        "tables": [],
        "error": "",
    }

    if not info["exists"]:
        return info

    try:
        file_size_bytes = os.path.getsize(DB_NAME)
        info["file_size_mb"] = round(file_size_bytes / 1024 / 1024, 3)

        modified_timestamp = os.path.getmtime(DB_NAME)
        info["modified_at"] = datetime.fromtimestamp(
            modified_timestamp
        ).strftime("%Y-%m-%d %H:%M:%S")

        conn = get_connection()
        cur = conn.cursor()

        cur.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
            ORDER BY name
            """
        )

        table_names = [row[0] for row in cur.fetchall()]

        for table_name in table_names:
            try:
                cur.execute(f'SELECT COUNT(*) FROM "{table_name}"')
                row_count = cur.fetchone()[0]
            except Exception:
                row_count = ""

            info["tables"].append(
                {
                    "테이블": table_name,
                    "데이터 건수": row_count,
                }
            )

        conn.close()

    except Exception as e:
        info["error"] = str(e)

    return info


def validate_sqlite_db_file(db_path):
    try:
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()

        cur.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
            """
        )

        tables = cur.fetchall()
        conn.close()

        if len(tables) == 0:
            return False, "업로드한 DB에 테이블이 없습니다."

        return True, ""

    except Exception as e:
        return False, str(e)


def render_db_management_page():
    st.subheader("DB 백업/복원")

    st.write(
        "현재 SQLite DB 파일을 백업하거나, 이전에 백업한 DB 파일로 복원합니다. "
        "복원 전에는 현재 DB가 자동으로 백업됩니다."
    )

    db_info = get_db_file_info()

    if not db_info["exists"]:
        st.error(f"DB 파일을 찾을 수 없습니다: {DB_NAME}")
        return

    st.divider()

    st.subheader("현재 DB 정보")

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("DB 파일명", DB_NAME)

    with col2:
        st.metric("파일 크기(MB)", f"{db_info['file_size_mb']:,.3f}")

    with col3:
        st.metric("마지막 수정일", db_info["modified_at"])

    if db_info["error"]:
        st.warning(f"DB 정보를 읽는 중 오류가 발생했습니다: {db_info['error']}")
    else:
        table_df = pd.DataFrame(db_info["tables"])

        if table_df.empty:
            st.info("DB에 테이블이 없습니다.")
        else:
            st.dataframe(
                table_df,
                use_container_width=True,
                hide_index=True,
            )

    st.divider()

    st.subheader("현재 DB 백업 다운로드")

    backup_file_name = f"ghg_manager_backup_{get_timestamp_text()}.db"

    with open(DB_NAME, "rb") as db_file:
        st.download_button(
            label="현재 DB 백업 다운로드",
            data=db_file,
            file_name=backup_file_name,
            mime="application/octet-stream",
        )

    st.caption(
        "다운로드한 .db 파일은 안전한 위치에 보관하세요. "
        "나중에 복원할 때 이 파일을 업로드하면 됩니다."
    )

    st.divider()

    st.subheader("백업 DB 업로드 후 복원")

    st.warning(
        "복원을 실행하면 현재 DB가 업로드한 DB로 교체됩니다. "
        "다만 복원 직전의 현재 DB는 db_backups 폴더에 자동 백업됩니다."
    )

    uploaded_db = st.file_uploader(
        "복원할 SQLite DB 파일을 업로드하세요.",
        type=["db", "sqlite", "sqlite3"],
    )

    restore_confirm_text = st.text_input(
        "복원을 진행하려면 아래 입력칸에 '복원'이라고 입력하세요.",
        value="",
    )

    if uploaded_db is not None:
        st.info(f"업로드한 파일명: {uploaded_db.name}")

        restore_button_disabled = restore_confirm_text.strip() != "복원"

        if st.button(
            "DB 복원 실행",
            type="primary",
            disabled=restore_button_disabled,
        ):
            try:
                os.makedirs("db_backups", exist_ok=True)

                timestamp_text = get_timestamp_text()

                before_restore_backup_path = os.path.join(
                    "db_backups",
                    f"backup_before_restore_{timestamp_text}.db"
                )

                uploaded_temp_path = os.path.join(
                    "db_backups",
                    f"uploaded_restore_{timestamp_text}.db"
                )

                # 1. 현재 DB를 먼저 자동 백업합니다.
                shutil.copy2(DB_NAME, before_restore_backup_path)

                # 2. 업로드한 DB 파일을 임시 위치에 저장합니다.
                with open(uploaded_temp_path, "wb") as temp_file:
                    temp_file.write(uploaded_db.getbuffer())

                # 3. 업로드한 파일이 SQLite DB인지 확인합니다.
                is_valid_db, validation_message = validate_sqlite_db_file(
                    uploaded_temp_path
                )

                if not is_valid_db:
                    st.error(
                        "업로드한 파일을 SQLite DB로 확인할 수 없습니다. "
                        f"오류 내용: {validation_message}"
                    )
                    st.info(
                        f"복원은 실행되지 않았습니다. 현재 DB 자동 백업 파일: {before_restore_backup_path}"
                    )
                    return

                # 4. 기존 DB를 업로드한 DB로 교체합니다.
                shutil.copy2(uploaded_temp_path, DB_NAME)

                st.success("DB 복원이 완료되었습니다.")
                st.info(f"복원 전 현재 DB 백업 파일: {before_restore_backup_path}")
                st.info("복원 결과를 완전히 반영하려면 브라우저를 새로고침하거나 Streamlit 앱을 재실행하세요.")

            except Exception as e:
                st.error(f"DB 복원 중 오류가 발생했습니다: {e}")


def main():
    st.set_page_config(
        page_title="온실가스 배출량 월별 관리 프로그램",
        layout="wide"
    )

    init_db()
    seed_default_factors()
    sync_activity_corporations()


    st.title("온실가스 배출량 월별 관리 프로그램")

    menu = st.sidebar.radio(
    "메뉴",
    [
        "대시보드",
        "온실가스 배출량 월별 DB",
        "사업장별 DB",
        "에너지원별 DB",
        "연간 월별 사용량 입력",
        "매출액 관리",
        "감축목표 관리",
        "입력 데이터 조회",
        "시설-에너지원 관리",
        "배출계수 관리",
        "DB 백업/복원",
    ]
)

    if menu == "대시보드":
        render_dashboard_page()

    elif menu == "온실가스 배출량 월별 DB":
        render_monthly_emission_db_page()

    elif menu == "사업장별 DB":
        render_site_db_page()    

    elif menu == "에너지원별 DB":
        render_energy_db_page()    

    elif menu == "연간 월별 사용량 입력":
        render_yearly_input_page()

    elif menu == "매출액 관리":
        render_revenue_page()    

    elif menu == "감축목표 관리":
        render_reduction_target_page()    

    elif menu == "입력 데이터 조회":
        render_data_edit_page()

    elif menu == "시설-에너지원 관리":
        render_facility_energy_page()

    elif menu == "배출계수 관리":
        render_factor_page()

    elif menu == "DB 백업/복원":
        render_db_management_page()
    

if __name__ == "__main__":
    main()

